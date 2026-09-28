"""Office 文档正文抽取：Word / PowerPoint / Excel → 带章节定位的纯文本。

与 PDF 解析同一个约定：只抽取可见文本，按文档自身结构（标题 / 幻灯片 / 工作表）
切成章节，每节带可读的定位（「第 3 张幻灯片」「工作表 Sheet1」），供逐字核验后
在报告里指回原文位置。解析失败抛 ``OfficeParseError``，不猜内容。
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

MAX_TABLE_ROWS = 2_000


class OfficeParseError(ValueError):
    pass


@dataclass(frozen=True)
class OfficeSection:
    title: str
    text: str


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sniff_office(raw: bytes, mime_type: str = "", filename: str = "") -> str | None:
    """识别 OOXML 容器类型：返回 docx / pptx / xlsx，或 None。"""
    name = filename.casefold()
    for kind, mime in (("docx", DOCX_MIME), ("pptx", PPTX_MIME), ("xlsx", XLSX_MIME)):
        if mime_type == mime or name.endswith(f".{kind}"):
            return kind
    if not raw.startswith(b"PK"):
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return None


def _docx_sections(raw: bytes) -> list[OfficeSection]:
    from docx import Document

    document = Document(io.BytesIO(raw))
    sections: list[OfficeSection] = []
    title = "正文"
    buffer: list[str] = []

    def flush() -> None:
        text = "\n\n".join(part for part in buffer if part.strip())
        if text.strip():
            sections.append(OfficeSection(title, text))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").casefold() if paragraph.style is not None else ""
        if style.startswith("heading") or style.startswith("标题") or style == "title":
            flush()
            title = text[:120]
        else:
            buffer.append(text)
    flush()
    for index, table in enumerate(document.tables, 1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows[:MAX_TABLE_ROWS]
        ]
        text = "\n".join(row for row in rows if row.strip(" |"))
        if text:
            sections.append(OfficeSection(f"表 {index}", text))
    return sections


def _pptx_sections(raw: bytes) -> list[OfficeSection]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw))
    sections: list[OfficeSection] = []
    for number, slide in enumerate(presentation.slides, 1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
        ]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(texts) + (f"\n\n演讲备注：{notes}" if notes else "")
        if body.strip():
            sections.append(OfficeSection(f"第 {number} 张幻灯片", body.strip()))
    return sections


def _xlsx_sections(raw: bytes) -> list[OfficeSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sections: list[OfficeSection] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True, max_row=MAX_TABLE_ROWS):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                sections.append(OfficeSection(f"工作表 {sheet.title}", "\n".join(rows)))
    finally:
        workbook.close()
    return sections


_PARSERS = {"docx": _docx_sections, "pptx": _pptx_sections, "xlsx": _xlsx_sections}
MIME_BY_KIND = {"docx": DOCX_MIME, "pptx": PPTX_MIME, "xlsx": XLSX_MIME}


def parse_office(raw: bytes, kind: str) -> list[OfficeSection]:
    """解析 docx / pptx / xlsx；任何库异常都转成 ``OfficeParseError``。"""
    parser = _PARSERS.get(kind)
    if parser is None:
        raise OfficeParseError(f"不支持的文档类型：{kind}")
    try:
        sections = parser(raw)
    except OfficeParseError:
        raise
    except Exception as exc:  # 损坏或加密的文档：给出可读原因，不外泄堆栈
        raise OfficeParseError(f"无法解析 {kind.upper()} 文档：{type(exc).__name__}") from exc
    if not sections:
        raise OfficeParseError(f"{kind.upper()} 文档中没有可提取的文字")
    return sections


__all__ = [
    "MIME_BY_KIND",
    "OfficeParseError",
    "OfficeSection",
    "parse_office",
    "sniff_office",
]
