"""块树 → DOCX（python-docx，统一的工作台报告样式）。

样式在这里一次定义（字体、字号、标题配色、表格边框、页边距、页码），所有报告类
DOCX 都从同一个 ``_base_document`` 开始——这是「报告有规定的样式基底，从基底改内容，
不另起一套」在本项目里的落点。
"""

from __future__ import annotations

import io
import re
from collections.abc import Mapping

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .markdown import Block, Inline, parse_blocks, plain
from .math import office_math

_BODY_FONT = "Microsoft YaHei"
_ACCENT = RGBColor(0x1F, 0x5F, 0x8B)
_MUTED = RGBColor(0x5B, 0x66, 0x75)
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _set_east_asian(run_or_style, name: str) -> None:  # type: ignore[no-untyped-def]
    element = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _page_number_footer(section) -> None:  # type: ignore[no-untyped-def]
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    run.font.size = Pt(8)
    run.font.color.rgb = _MUTED


def _base_document():  # type: ignore[no-untyped-def]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.4))
    section.top_margin = section.bottom_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(10.5)
    _set_east_asian(normal, _BODY_FONT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)
    for level, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        style = document.styles[f"Heading {level}"]
        style.font.name = _BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _ACCENT if level > 1 else RGBColor(0x10, 0x28, 0x3D)
        _set_east_asian(style, _BODY_FONT)
        style.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        style.paragraph_format.space_after = Pt(6)
    _page_number_footer(section)
    return document


def _add_runs(paragraph, inlines: list[Inline]) -> None:  # type: ignore[no-untyped-def]
    for item in inlines:
        if item.math:
            _add_math(paragraph, item.text, item.display)
            continue
        run = paragraph.add_run(item.text)
        run.bold = item.bold or None
        run.italic = item.italic or item.math or None
        if item.code or item.math:
            run.font.name = "Consolas" if item.code else "Cambria Math"
            _set_east_asian(run, run.font.name)


def _add_math(paragraph, tex: str, display: bool = False) -> None:  # type: ignore[no-untyped-def]
    from lxml import etree

    xml = office_math(tex, display)
    xml = xml.replace(
        "<m:oMath>",
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
        1,
    )
    root = etree.fromstring(xml.encode())
    for run in root.findall(".//" + qn("m:r")):
        properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Cambria Math")
        fonts.set(qn("w:hAnsi"), "Cambria Math")
        fonts.set(qn("w:eastAsia"), _BODY_FONT)
        properties.append(fonts)
        run.insert(1 if run.find(qn("m:rPr")) is not None else 0, properties)
    if display:
        wrapper = OxmlElement("m:oMathPara")
        wrapper.append(root)
        paragraph._p.append(wrapper)
    else:
        paragraph._p.append(root)


def _shade(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _table(document, block: Block) -> None:  # type: ignore[no-untyped-def]
    width = max(len(row) for row in block.rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(block.rows):
        cells = table.add_row().cells
        for col in range(width):
            paragraph = cells[col].paragraphs[0]
            _add_runs(paragraph, row[col] if col < len(row) else [])
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                _shade(cells[col], "EEF2F6")
    document.add_paragraph()


def render_docx(
    markdown: str,
    *,
    title: str,
    meta: str = "",
    images: Mapping[str, bytes] | None = None,
) -> bytes:
    images = images or {}
    # python-docx 拒收 XML 1.0 不允许的控制字符（网页快照里常见），整份 Word 会丢
    markdown, title, meta = (_XML_ILLEGAL.sub("", value) for value in (markdown, title, meta))
    blocks = parse_blocks(markdown)
    document = _base_document()
    title_inlines = [Inline(title)]
    if blocks and blocks[0].kind == "heading" and blocks[0].level == 1:
        title_inlines = blocks[0].inlines
        title = plain(blocks[0].inlines) or title
        blocks = blocks[1:]
    document.core_properties.title = title[:200]
    _add_runs(document.add_heading(level=1), title_inlines)
    if meta:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(meta)
        run.font.size = Pt(9)
        run.font.color.rgb = _MUTED
    for block in blocks:
        if block.kind == "heading":
            _add_runs(document.add_heading(level=min(max(block.level, 2), 4)), block.inlines)
        elif block.kind == "math":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_math(paragraph, block.text, True)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            if len(block.inlines) == 1 and block.inlines[0].math:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(paragraph, block.inlines)
        elif block.kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            _add_runs(paragraph, block.inlines)
            for run in paragraph.runs:
                run.font.color.rgb = _MUTED
        elif block.kind == "list":
            for item in block.items:
                style = "List Number" if item.ordered else "List Bullet"
                if item.depth:
                    style += f" {min(item.depth + 1, 3)}"
                paragraph = document.add_paragraph(style=style)
                _add_runs(paragraph, item.inlines)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif block.kind == "table" and block.rows:
            _table(document, block)
        elif block.kind == "image":
            data = images.get(block.src) or images.get(block.src.rsplit("/", 1)[-1])
            if data is None:
                document.add_paragraph(f"［图缺失：{block.text}］")
                continue
            document.add_picture(io.BytesIO(data), width=Cm(15))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = document.add_paragraph(block.text)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = _MUTED
        elif block.kind == "rule":
            document.add_paragraph("―" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_stats(data: bytes) -> dict[str, int]:
    """验收门用：DOCX 里的标题数、表格数、内嵌图片数与正文字符数。"""
    document = Document(io.BytesIO(data))
    headings = sum(
        1
        for p in document.paragraphs
        if p.style is not None and (p.style.name or "").startswith("Heading")
    )
    images = len(document.inline_shapes)
    text = "".join(p.text for p in document.paragraphs)
    return {
        "headings": headings,
        "tables": len(document.tables),
        "images": images,
        "chars": len(text),
    }


__all__ = ["docx_stats", "render_docx"]
