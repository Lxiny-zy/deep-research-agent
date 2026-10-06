"""Native PowerPoint visuals from frozen, reviewed content; no fetching or inference."""

from __future__ import annotations

import hashlib
import io
import math
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt


class LiteralCategoryChartData(CategoryChartData):
    """Keep source labels literal in the editable embedded spreadsheet."""

    @property
    def xlsx_blob(self) -> bytes:
        from xlsxwriter import Workbook

        output = io.BytesIO()
        with Workbook(
            output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
        ) as workbook:
            sheet = workbook.add_worksheet()
            for row, category in enumerate(self.categories, 1):
                sheet.write_string(row, 0, str(category.label))
            for column, series in enumerate(self, 1):
                sheet.write_string(0, column, str(series.name))
                for row, value in enumerate(series.values, 1):
                    sheet.write_number(row, column, value)
        return output.getvalue()


def draw_visual(slide: Any, data: dict[str, Any], images: dict[str, bytes], *, top: int) -> None:
    from .pptx import _INK, _MUTED, _SOFT, _estimated_lines, _text

    caption = "\n".join(filter(None, [data.get("caption", ""), *data.get("scope_notes", [])]))
    if data["kind"] == "image":
        caption += "\n图片可移动/缩放，底层内容不可编辑。"
    caption_lines = _estimated_lines(caption, width=160, reference=True) if caption else 0
    if caption_lines > 5:
        raise ValueError("图表口径说明超过可读页高，请拆分内容，不能截去必要条件")
    caption_height = max(300000, int(Pt(caption_lines * 14)))
    bottom = 6100000 - caption_height
    width, height = 10800000, bottom - top - 100000
    kind = data["kind"]
    if kind == "table":
        rows = [data["headers"], *data["rows"]]
        if not rows or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("幻灯片表格行列不完整")
        cell_width = max(12, 120 // len(rows[0]))
        needed_height = 0
        for row in rows:
            row_lines = max(
                _estimated_lines(str(cell), width=cell_width, reference=True) for cell in row
            )
            if row_lines > 3:
                raise ValueError("表格单元格过长，请缩减列数或拆分说明，不能缩成不可读字号")
            needed_height += int(Pt(row_lines * 17 + 8))
        if needed_height > height:
            raise ValueError("表格行高与口径说明超出页面，请按问题拆分表格，不能裁切内容")
        shape = slide.shapes.add_table(
            len(rows), len(rows[0]), Emu(700000), Emu(top), Emu(width), Emu(height)
        )
        shape.name = "DR-table-" + data["id"]
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                cell = shape.table.cell(row_index, column_index)
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    _SOFT if row_index == 0 else _INK.__class__(255, 255, 255)
                )
                _text(cell.text_frame, str(value), size=14, bold=row_index == 0, color=_INK)
    elif kind in {"bar", "line"}:
        if len(data["categories"]) > 10 or any(
            _estimated_lines(label, width=24, reference=True) > 2
            for label in data["categories"]
        ):
            raise ValueError("图表分类过多或标签过长，请使用原生表格保留完整对象名称")
        chart_data = LiteralCategoryChartData()
        chart_data.categories = data["categories"]
        for series in data["series"]:
            if len(series["values"]) != len(data["categories"]) or any(
                isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)
                for v in series["values"]
            ):
                raise ValueError("原生图表数值缺失或无效，不能自动填零")
            chart_data.add_series(series["name"], series["values"])
        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED if kind == "bar" else XL_CHART_TYPE.LINE_MARKERS,
            Emu(700000),
            Emu(top),
            Emu(width),
            Emu(height),
            chart_data,
        )
        shape.name = "DR-chart-" + data["id"]
        chart = shape.chart
        chart.has_legend = len(data["series"]) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.category_axis.tick_labels.font.size = Pt(12)
        chart.value_axis.tick_labels.font.size = Pt(12)
        chart.value_axis.has_title = bool(data.get("unit"))
        if data.get("unit"):
            _text(chart.value_axis.axis_title.text_frame, data["unit"], size=12)
        if kind == "bar" and all(v >= 0 for series in data["series"] for v in series["values"]):
            chart.value_axis.minimum_scale = 0
        chart.plots[0].has_data_labels = True
        chart.plots[0].data_labels.font.size = Pt(11)
        for series in chart.series:
            if kind == "line":
                series.smooth = False
    elif kind == "image":
        asset = data.get("asset", "")
        blob = images.get(asset)
        if not blob or len(blob) > 16 * 1024 * 1024 or not data.get("citations"):
            raise ValueError("图片必须来自已冻结且有引用的素材，不能临时下载或猜测来源")
        if hashlib.sha256(blob).hexdigest() != data.get("sha256"):
            raise ValueError("图片与冻结素材哈希不一致")
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as probe:
            w, h = probe.size
            if w * h > 30_000_000:
                raise ValueError("图片像素超过单页处理上限")
        factor = min(width / w, height / h)
        shown_width, shown_height = int(w * factor), int(h * factor)
        shape = slide.shapes.add_picture(
            io.BytesIO(blob),
            Emu(700000 + (width - shown_width) // 2),
            Emu(top + (height - shown_height) // 2),
            Emu(shown_width),
            Emu(shown_height),
        )
        shape.name = "DR-image-" + data["id"]
    else:
        raise ValueError("不支持的幻灯片图形，不能静默省略")
    if caption:
        box = slide.shapes.add_textbox(
            Emu(700000), Emu(bottom + 50000), Emu(width), Emu(caption_height)
        )
        _text(box.text_frame, caption, size=12, color=_MUTED)


def visual_consistency(pptx: bytes, deck: dict[str, Any], markdown: str | None = None) -> list[str]:
    """Check actual objects and embedded values, not just a self-reported schema."""
    from .markdown import parse_blocks
    from .pptx import paginate_deck

    def visible_tables(text: str) -> list[str]:
        return [block.plain() for block in parse_blocks(text) if block.kind == "table"]

    expected = [s["visual_data"] for s in paginate_deck(deck)["slides"] if s.get("visual_data")]
    presentation = Presentation(io.BytesIO(pptx))
    actual = [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.name.startswith("DR-")
    ]
    if len(actual) != len(expected):
        return ["PPTX 图表/图片数量与冻结结构不一致"]
    issues = []
    for data, shape in zip(expected, actual, strict=True):
        label = data.get("title") or data["id"]
        kind = data["kind"]
        try:
            if kind == "table":
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if rows != [data["headers"], *data["rows"]]:
                    issues.append(f"PPTX 表格 {label} 的单元格与冻结数据不一致")
            elif kind in {"bar", "line"}:
                chart = shape.chart
                if [str(c.label) for c in chart.plots[0].categories] != data["categories"]:
                    issues.append(f"PPTX 图表 {label} 的分类标签不一致")
                if len(chart.series) != len(data["series"]) or any(
                    series.name != spec["name"]
                    or len(series.values) != len(spec["values"])
                    or any(
                        not math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
                        for a, b in zip(series.values, spec["values"], strict=True)
                    )
                    for series, spec in zip(chart.series, data["series"], strict=True)
                ):
                    issues.append(f"PPTX 图表 {label} 的数值或系列名称不一致")
            elif kind == "image" and hashlib.sha256(shape.image.blob).hexdigest() != data.get(
                "sha256"
            ):
                issues.append(f"PPTX 图片 {label} 与冻结素材不一致")
            if kind != "image" and markdown is not None:
                if any(
                    table not in visible_tables(markdown)
                    for table in visible_tables(data.get("markdown", ""))
                ):
                    issues.append(f"PPTX 图表 {label} 缺少一致的 Markdown 数据表")
        except (AttributeError, KeyError, TypeError, ValueError):
            issues.append(f"PPTX 图形 {label} 的原生对象无法核验")
    return issues
