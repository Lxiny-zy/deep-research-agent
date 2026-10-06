"""SlideDeck → PPTX（python-pptx，16:9，统一版式，每页附演讲备注）。

版式由代码控制，模型提供内容。行数和时长均为估算；实际应用渲染和试讲仍须验收。

``fit_report`` 是版面自检：每条要点的估算行数、每页总行数是否超出文本框容量。
验收门据此报告溢出页，而不是等用户打开才发现。
"""

from __future__ import annotations

import io
import math
import re
from collections.abc import Iterator
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from ..slide_content import meaningful_notes
from .markdown import _inlines, _parser
from .math import MathRenderError, math_asset, office_math

_W, _H = Emu(12192000), Emu(6858000)  # 16:9
_INK = RGBColor(0x14, 0x22, 0x30)
_ACCENT = RGBColor(0x1F, 0x5F, 0x8B)
_SOFT = RGBColor(0xEE, 0xF3, 0xF7)
_MUTED = RGBColor(0x5B, 0x66, 0x75)
_FONT = "Microsoft YaHei"
_BULLET_PT = 20
_MAX_LINES = 11
_CHARS_PER_LINE = 34  # 20pt 中文在 11.2 英寸宽文本框中的近似容量
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_DRAWING_2010_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
_COVER_LABELS = {"标题页", "封面", "title", "titleslide"}


def _font(font: Any, name: str) -> None:
    font.name = name
    east_asian = font._rPr.find(f"{{{_DRAWING_NS}}}ea")
    if east_asian is None:
        east_asian = etree.SubElement(font._rPr, f"{{{_DRAWING_NS}}}ea")
    east_asian.set("typeface", _FONT)


def _paragraph(
    paragraph: Any, text: str, *, size: int, bold: bool = False, color: RGBColor = _INK
) -> None:
    paragraph.font.size = Pt(size)
    _font(paragraph.font, _FONT)
    for item in _inlines(_parser().parseInline(text)[0]):
        if item.math:
            # PowerPoint's Office 2010 DrawingML math extension carries editable
            # OMML, rather than displaying TeX delimiters and commands as text.
            xml = office_math(item.text, item.display)
            root = etree.fromstring(
                f'<a14:m xmlns:a14="{_DRAWING_2010_NS}" xmlns:m="{_MATH_NS}" '
                f'xmlns:a="{_DRAWING_NS}"><m:oMathPara>{xml}</m:oMathPara></a14:m>'.encode()
            )
            if item.text.lstrip().startswith(("_", "^")):
                # TeX permits a script with an empty base after plain text
                # (head$_j$). Office shows an editable placeholder for an empty
                # base; a zero-width text base retains the intended appearance.
                for empty in root.findall(f".//{{{_MATH_NS}}}e"):
                    if not "".join(empty.itertext()).strip():
                        text_node = empty.find(f".//{{{_MATH_NS}}}t")
                        if text_node is None:
                            run = etree.SubElement(empty, f"{{{_MATH_NS}}}r")
                            text_node = etree.SubElement(run, f"{{{_MATH_NS}}}t")
                        text_node.text = "\u200b"
            for node in root.findall(f".//{{{_MATH_NS}}}r"):
                properties = etree.Element(f"{{{_DRAWING_NS}}}rPr", sz=str(size * 100))
                fill = etree.SubElement(properties, f"{{{_DRAWING_NS}}}solidFill")
                etree.SubElement(fill, f"{{{_DRAWING_NS}}}srgbClr", val=str(color))
                etree.SubElement(properties, f"{{{_DRAWING_NS}}}latin", typeface="Cambria Math")
                etree.SubElement(properties, f"{{{_DRAWING_NS}}}ea", typeface=_FONT)
                node.insert(1 if node.find(f"{{{_MATH_NS}}}rPr") is not None else 0, properties)
            paragraph._p.append(root)
        else:
            run = paragraph.add_run()
            run.text = item.text
            run.font.size = Pt(size)
            run.font.bold = bold or item.bold
            run.font.italic = item.italic
            _font(run.font, "Consolas" if item.code else _FONT)
            run.font.color.rgb = color


def _text(frame, text: str, *, size: int, bold: bool = False, color: RGBColor = _INK) -> None:  # type: ignore[no-untyped-def]
    frame.clear()
    frame.word_wrap = True
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    _paragraph(paragraph, text, size=size, bold=bold, color=color)


def _band(slide, top: int, height: int, color: RGBColor) -> None:  # type: ignore[no-untyped-def]
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, _W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _weight(char: str, reference: bool = False) -> int:
    return 2 if ord(char) > 0x2E80 or (reference and char in "MW@") else 1


def _math_spans(text: str) -> dict[int, tuple[int, str, bool]]:
    """Use the renderer's actual math rules to preserve raw source coordinates."""
    md = _parser()
    spans = {}

    def capture(rule):  # type: ignore[no-untyped-def]
        def wrapped(state, silent):  # type: ignore[no-untyped-def]
            start, count = state.pos, len(state.tokens)
            matched = rule(state, silent)
            if matched and not silent:
                for token in state.tokens[count:]:
                    if token.type in {"math_inline", "math_inline_double"}:
                        spans[start] = (
                            state.pos,
                            token.content,
                            token.type == "math_inline_double",
                        )
            return matched

        return wrapped

    for name, rule in zip(
        md.inline.ruler.get_active_rules(), md.inline.ruler.getRules(""), strict=True
    ):
        if name in {"math_inline", "bracket_math"}:
            md.inline.ruler.at(name, capture(rule))
    md.inline.parse(text, md, {}, [])
    return spans


def _atoms(text: str, reference: bool = False) -> Iterator[tuple[int, int, int, float]]:
    spans = _math_spans(text)
    position = 0
    size = 12 if reference else _BULLET_PT
    while position < len(text):
        if position in spans:
            end, tex, display = spans[position]
            asset = math_asset(tex, display, size)
            yield (
                position,
                end,
                math.ceil(asset.width / (size / 2)),
                max(1.0, asset.height / (size * 1.2)),
            )
            position = end
        else:
            yield position, position + 1, _weight(text[position], reference), 1
            position += 1


def _estimated_lines(
    text: str, *, width: int = _CHARS_PER_LINE * 2, reference: bool = False
) -> int:
    total, used, height = 0.0, 0, 1.0
    for start, end, weight, atom_height in _atoms(text, reference):
        if weight > width:
            raise MathRenderError("公式超过幻灯片行宽，请拆分表达式")
        if text[start:end] == "\n" or used + weight > width:
            total += height
            used, height = 0, 1
        if text[start:end] != "\n":
            used += weight
            height = max(height, atom_height)
    return math.ceil(total + height)


def fit_report(deck: dict[str, Any]) -> list[dict[str, Any]]:
    """估算每页要点占用的行数，返回超出容量的页面。"""
    problems: list[dict[str, Any]] = []
    for number, slide in enumerate(deck.get("slides", []), 1 if _has_title_slide(deck) else 2):
        bullets = slide.get("bullets", [])
        heading_lines = _estimated_lines(str(slide.get("title", "")), width=60)
        if heading_lines > 2:
            problems.append({"slide": number, "lines": heading_lines,
                             "bullets": len(bullets), "title": True})
        visual = slide.get("visual_data") or slide.get("visual")
        limit = 3 if visual else _MAX_LINES
        lines = sum(_estimated_lines(bullet) for bullet in bullets)
        if lines > limit or len(bullets) > (2 if visual else 5):
            problems.append({"slide": number, "lines": lines, "bullets": len(bullets)})
        if not bullets and not visual:
            problems.append({"slide": number, "lines": 0, "bullets": 0, "empty": True})
    return problems


def _has_title_slide(deck: dict[str, Any]) -> bool:
    def title(value: str) -> str:
        return "".join(
            "".join(item.text for item in _inlines(_parser().parseInline(value)[0]))
            .casefold()
            .split()
        )

    slides = deck.get("slides", [])
    heading = title(str(deck.get("title", "")))
    first = title(str(slides[0].get("title", ""))) if slides else ""
    return bool(heading and first and (first == heading or first in _COVER_LABELS))


def _split_bullet(
    text: str, *, columns: int = _CHARS_PER_LINE * 2, max_lines: int = 4, reference: bool = False
) -> list[str]:
    """Bound individual paragraphs without throwing away their remaining text."""
    pieces: list[str] = []
    while text:
        width = 0
        lines, height = 0.0, 1.0
        end = 0
        protected = _math_spans(text)
        for start, stop, weight, atom_height in _atoms(text, reference):
            if weight > columns:
                raise MathRenderError("公式超过幻灯片行宽，请拆分表达式")
            newline = text[start:stop] == "\n"
            if newline or width + weight > columns:
                if lines + height >= max_lines:
                    break
                lines += height
                width, height = 0, 1
            if not newline:
                if end and lines + max(height, atom_height) > max_lines:
                    break
                width += weight
                height = max(height, atom_height)
            end = stop
        if end < len(text):
            # Prefer a sentence/word boundary, while keeping every character.
            boundary = max(text.rfind(mark, 0, end) for mark in (" ", "。", "；", ";"))
            if boundary >= end // 2 and not any(
                a < boundary + 1 < b for a, (b, _, _) in protected.items()
            ):
                end = boundary + 1
            opening = text.rfind("[", 0, end)
            if (
                opening > text.rfind("]", 0, end)
                and opening > 0
                and not any(a <= opening < b for a, (b, _, _) in protected.items())
            ):
                end = opening
        pieces.append(text[:end])
        text = text[end:]
    return pieces


def reference_pages(references: list[str]) -> list[list[str]]:
    """Fit full citations at readable size; long entries continue without clipping."""
    pages: list[list[str]] = []
    page: list[str] = []
    height = 0
    for index, reference in enumerate(references, 1):
        pieces = _split_bullet(reference, columns=100, max_lines=18, reference=True) or [""]
        for part, piece in enumerate(pieces):
            text = f"[{index}] " + ("（续）" if part else "") + piece
            needed = _estimated_lines(text, width=100, reference=True) * 16 + 8
            if page and height + needed > 360:
                pages.append(page)
                page, height = [], 0
            page.append(text)
            height += needed
    if page:
        pages.append(page)
    return pages


def paginate_deck(deck: dict[str, Any]) -> dict[str, Any]:
    """Continue overflowing content on another slide, preserving notes/citations."""
    if deck.get("_paginated"):
        return deck
    slides: list[dict[str, Any]] = []
    for source_index, spec in enumerate(deck.get("slides", [])):
        visual = spec.get("visual_data")
        visual_rows = visual.get("rows", []) if visual and visual["kind"] == "table" else []
        limit = 3 if visual or spec.get("visual") else _MAX_LINES
        pages: list[list[str]] = [[]]
        lines = 0
        for bullet in spec.get("bullets", []):
            for piece in _split_bullet(str(bullet), max_lines=min(4, limit)):
                needed = _estimated_lines(piece)
                if pages[-1] and (len(pages[-1]) >= 5 or lines + needed > limit):
                    pages.append([])
                    lines = 0
                pages[-1].append(piece)
                lines += needed
        # Avoid a dense page followed by a one-item continuation when the same
        # ordered content can be distributed more evenly without adding pages.
        for index in range(len(pages) - 2, -1, -1) if not visual else []:
            left, right = pages[index], pages[index + 1]
            if len(right) != 1:
                continue
            while len(left) > 1 and len(right) < 5:
                left_lines = sum(_estimated_lines(text) for text in left)
                right_lines = sum(_estimated_lines(text) for text in right)
                moved = _estimated_lines(left[-1])
                if right_lines + moved > _MAX_LINES or abs(
                    left_lines - moved - right_lines - moved
                ) >= abs(left_lines - right_lines):
                    break
                right.insert(0, left.pop())
        if visual_rows:
            while len(pages) < math.ceil(len(visual_rows) / 6):
                pages.append([])
        # Keep sentence order and every character, but never copy a whole script
        # to every continuation. Insufficient scripts remain visibly incomplete.
        chunks = re.findall(
            r".+?(?:[。！？!?]+|\.(?=\s|$)|\n|$)", str(spec.get("notes", "")), re.S
        )
        notes = [""] * len(pages)
        for index, chunk in enumerate(chunks):
            notes[min(len(pages) - 1, index * len(pages) // max(1, len(chunks)))] += chunk
        for index, bullets in enumerate(pages):
            page_visual = None
            if visual and visual_rows and index * 6 < len(visual_rows):
                page_visual = {**visual, "rows": visual_rows[index * 6 : (index + 1) * 6]}
            elif visual and not visual_rows and index == 0:
                page_visual = visual
            slides.append(
                {
                    **spec,
                    "bullets": bullets,
                    "notes": notes[index],
                    "visual_data": page_visual,
                    "_source_index": source_index,
                    "title": str(spec.get("title", "")) + (f"（续 {index + 1}）" if index else ""),
                }
            )
    return {**deck, "slides": slides, "_paginated": True}


def render_pptx(
    deck: dict[str, Any],
    *,
    citations: list[str] | None = None,
    images: dict[str, bytes] | None = None,
) -> bytes:
    deck = paginate_deck(deck)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = _W, _H
    blank = presentation.slide_layouts[6]

    content_cover = _has_title_slide(deck)
    if not content_cover:
        cover = presentation.slides.add_slide(blank)
        _band(cover, 0, _H, _ACCENT)
        box = cover.shapes.add_textbox(Emu(800000), Emu(2200000), Emu(10600000), Emu(1400000))
        _text(
            box.text_frame,
            deck.get("title", "研究汇报"),
            size=40,
            bold=True,
            color=RGBColor(255, 255, 255),
        )
        if deck.get("subtitle"):
            sub = cover.shapes.add_textbox(Emu(800000), Emu(3700000), Emu(10600000), Emu(800000))
            _text(sub.text_frame, deck["subtitle"], size=20, color=RGBColor(0xDD, 0xE8, 0xF1))
        cover.notes_slide.notes_text_frame.text = "开场：介绍题目与汇报结构。"

    slides = deck.get("slides", [])
    for number, spec in enumerate(slides, 1 if content_cover else 2):
        slide = presentation.slides.add_slide(blank)
        _band(slide, 0, Emu(1150000), _SOFT)
        _band(slide, Emu(1150000), Emu(40000), _ACCENT)
        title = slide.shapes.add_textbox(Emu(600000), Emu(250000), Emu(11000000), Emu(750000))
        heading = str(spec.get("title", ""))
        if content_cover and number == 1 and "".join(heading.casefold().split()) in _COVER_LABELS:
            heading = str(deck.get("title", heading))
        _text(title.text_frame, heading, size=28, bold=True)
        title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        title.text_frame.paragraphs[0].line_spacing = 1.0
        body_top = 1450000
        if content_cover and number == 1 and deck.get("subtitle"):
            subtitle = str(deck["subtitle"])
            subtitle_height = int(Pt(_estimated_lines(subtitle, width=140, reference=True) * 15))
            sub = slide.shapes.add_textbox(
                Emu(700000), Emu(1200000), Emu(10800000), Emu(subtitle_height)
            )
            _text(sub.text_frame, subtitle, size=12, color=_MUTED)
            body_top = max(body_top, 1200000 + subtitle_height + 80000)
        visual = spec.get("visual_data")
        body = slide.shapes.add_textbox(
            Emu(700000),
            Emu(body_top),
            Emu(10800000),
            Emu(1200000 if visual else 6050000 - body_top),
        )
        frame = body.text_frame
        frame.word_wrap = True
        frame.clear()
        for index, bullet in enumerate(spec.get("bullets", [])):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            _paragraph(paragraph, f"•  {bullet}", size=_BULLET_PT)
            paragraph.space_after = Pt(10)
        if visual:
            from .pptx_visuals import draw_visual

            draw_visual(
                slide, visual, images or {}, top=2750000 if spec.get("bullets") else 1400000
            )
        cites = [int(i) for i in spec.get("citations", []) if isinstance(i, int)]
        if cites:
            foot = slide.shapes.add_textbox(Emu(700000), Emu(6250000), Emu(10800000), Emu(400000))
            _text(
                foot.text_frame, "来源：" + " ".join(f"[{i}]" for i in cites), size=11, color=_MUTED
            )
        page = slide.shapes.add_textbox(Emu(11000000), Emu(6300000), Emu(900000), Emu(350000))
        _text(page.text_frame, str(number), size=11, color=_MUTED)
        page.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        note = str(spec.get("notes", ""))
        slide.notes_slide.notes_text_frame.text = note if meaningful_notes(note) else ""

    for page_index, references in enumerate(reference_pages(citations or [])):
        refs = presentation.slides.add_slide(blank)
        _band(refs, 0, Emu(1150000), _SOFT)
        title = refs.shapes.add_textbox(Emu(600000), Emu(250000), Emu(11000000), Emu(750000))
        heading = "参考文献" if page_index == 0 else "参考文献（续）"
        _text(title.text_frame, heading, size=28, bold=True)
        body = refs.shapes.add_textbox(Emu(700000), Emu(1450000), Emu(10800000), Emu(4800000))
        frame = body.text_frame
        frame.word_wrap = True
        frame.margin_top = frame.margin_bottom = 0
        frame.clear()
        for offset, reference in enumerate(references):
            paragraph = frame.paragraphs[0] if offset == 0 else frame.add_paragraph()
            _paragraph(paragraph, reference, size=12)
            paragraph.line_spacing = Pt(16)
            paragraph.space_after = Pt(8)
        refs.notes_slide.notes_text_frame.text = "参考文献列表。"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def pptx_stats(data: bytes) -> dict[str, int]:
    presentation = Presentation(io.BytesIO(data))
    notes = sum(
        1
        for slide in presentation.slides
        if slide.has_notes_slide and meaningful_notes(slide.notes_slide.notes_text_frame.text)
    )
    return {"slides": len(presentation.slides), "with_notes": notes}


__all__ = ["fit_report", "pptx_stats", "render_pptx"]
