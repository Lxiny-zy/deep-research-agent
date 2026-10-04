"""Bind a task's terminal state to its frozen delivery and quality evidence."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from zipfile import ZipFile

from ..blocking import run_blocking, run_rendering
from ..config import Settings
from ..persistence.repository import RUN_TERMINAL_STATUSES, ResearchRepository, RunDetail
from .contract import contract_from_scratch
from .delivery_store import build_or_load, current_version, load_version
from .gates import GateResult
from .publish import DeliveryBundle, build_bundle, delivery_fingerprint, resolve_template

COMPLETION_KEY = "_completion"


def file_issues(bundle: DeliveryBundle) -> list[str]:
    """Open every promised format, including formats outside prose consistency."""
    issues: list[str] = []
    for file in bundle.files:
        try:
            if not file.data:
                raise ValueError("empty file")
            if file.format == "png":
                from PIL import Image

                with Image.open(BytesIO(file.data)) as image:
                    if image.format != "PNG":
                        raise ValueError("wrong image format")
                    image.verify()
            elif file.format == "xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(BytesIO(file.data), read_only=True)
                try:
                    if not workbook.sheetnames:
                        raise ValueError("empty workbook")
                    for sheet in workbook:
                        for _ in sheet.iter_rows(values_only=True):
                            pass
                finally:
                    workbook.close()
            elif file.format in {"docx", "pptx"}:
                from xml.etree import ElementTree

                with ZipFile(BytesIO(file.data)) as archive:
                    if archive.testzip() is not None:
                        raise ValueError("corrupt archive")
                    main = "word/document.xml" if file.format == "docx" else "ppt/presentation.xml"
                    ElementTree.fromstring(archive.read(main))
                    ElementTree.fromstring(archive.read("[Content_Types].xml"))
                if file.format == "docx":
                    from docx import Document

                    Document(BytesIO(file.data))
                else:
                    from pptx import Presentation

                    Presentation(BytesIO(file.data))
            elif file.format == "pdf":
                import pymupdf

                with pymupdf.open(stream=file.data, filetype="pdf") as pdf:
                    if not len(pdf) or not any(page.get_text().strip() for page in pdf):
                        raise ValueError("no readable pages")
            elif file.format in {"md", "html", "json"}:
                text = file.data.decode("utf-8")
                if not text.strip():
                    raise ValueError("empty text")
                if file.format == "json":
                    import json

                    json.loads(text)
            else:
                raise ValueError("unvalidated format")
        except Exception:
            issues.append(f"{file.name} 无法通过文件格式与可读取性检查")
    return issues


def validate_bundle_files(bundle: DeliveryBundle) -> None:
    """Record readability failures before publishing, so format retry can repair them."""
    failures: list[str] = []
    failed_formats: set[str] = set()
    for file in bundle.files:
        isolated = DeliveryBundle(
            bundle.template, bundle.title, [file], [], "pass", bundle.generated_at
        )
        problems = file_issues(isolated)
        if not problems:
            continue
        failures.extend(problems)
        failed_formats.add(file.format)
        file.status = "fail"
        file.issues = list(dict.fromkeys([*file.issues, *problems]))
        if not any(item["format"] == file.format for item in bundle.failures):
            bundle.failures.append(
                {"format": file.format, "title": file.title, "issues": problems, "retryable": True}
            )
    bundle.gates = [gate for gate in bundle.gates if gate.name != "file_readability"]
    bundle.gates.append(
        GateResult(
            "file_readability",
            "fail" if failures else "pass",
            failures,
            {"failed_formats": sorted(failed_formats)},
        )
    )
    if failures:
        bundle.status = "fail"


def promised_formats(detail: RunDetail) -> set[str] | None:
    scratch = detail.orchestration.checkpoint.get("scratch", {}) if detail.orchestration else {}
    contract = contract_from_scratch(scratch)
    workbench = scratch.get("workbench", {})
    # A historical/custom workflow must not inherit autoResearch's promises
    # merely because that is the delivery panel's fallback template.
    if contract is None and not (isinstance(workbench, dict) and workbench.get("template")):
        return None
    wants = (
        contract.deliverables
        if contract and contract.deliverables
        else resolve_template(detail).deliverables
    )
    expanded: set[str] = set()
    for item in wants:
        expanded.update(
            {"html", "png"} if item == "mindmap" else {"xlsx"} if item == "statistics" else {item}
        )
    return expanded


def assess_completion(detail: RunDetail, bundle: DeliveryBundle) -> dict[str, Any]:
    required = promised_formats(detail)
    if required is None:
        raise ValueError("task has no frozen delivery contract")
    issues = file_issues(bundle)
    template = resolve_template(detail)
    required_gates = {"markdown", "structure", "length", "consistency", "file_readability"}
    required_gates.add("node_evidence" if template.key == "mindmap" else "prose_evidence")
    if template.key == "dataAnalysis":
        required_gates.add("analysis")
    if template.key == "peerReview":
        required_gates.add("review")
    actual_gates = {gate.name for gate in bundle.gates}
    if required_gates - actual_gates:
        issues.append("缺少必需验收记录：" + "、".join(sorted(required_gates - actual_gates)))
    for gate in bundle.gates:
        if gate.status != "pass":
            issues.extend(gate.issues or [f"交付检查 {gate.name} 尚未通过"])
        if gate.name == "prose_evidence" and gate.metrics.get("method") == "not_reviewed":
            issues.append("最终正文尚未完成模型支持关系核验")
    for failure in bundle.failures:
        issues.extend(failure.get("issues") or [f"{failure.get('format', '文件')} 尚未生成"])
    for fmt in sorted(required):
        files = [file for file in bundle.files if file.format == fmt]
        if not files:
            issues.append(f"缺少承诺的 {fmt.upper()} 交付文件")
        elif any(file.status != "pass" or not file.data for file in files):
            issues.append(f"{fmt.upper()} 交付文件尚未通过验收")
    if bundle.status != "pass" and not issues:
        issues.append("交付整体检查尚未通过")
    source = delivery_fingerprint(detail)
    if bundle.input_version != source:
        issues.append("交付文件与当前定稿版本不一致")
    return {
        "policy_version": 1,
        "status": "needs_review" if issues else "done",
        "scope": "frozen_task_delivery",
        "input_version": source,
        "content_version": bundle.content_version,
        "required_formats": sorted(required),
        "required_gates": sorted(required_gates),
        "issues": list(dict.fromkeys(issues)),
        "gates": [gate.to_dict() for gate in bundle.gates],
        "checked_at": datetime.now(UTC).isoformat(),
    }


async def prepare_completion(detail: RunDetail, settings: Settings) -> dict[str, Any] | None:
    if promised_formats(detail) is None:
        return None
    bundle = await run_rendering(
        build_or_load, detail, settings.artifact_root, settings.artifact_total_bytes, build_bundle
    )
    return await run_rendering(assess_completion, detail, bundle)


async def synchronize_completion(repo: ResearchRepository, run_id: str, settings: Settings) -> bool:
    """Reconcile a committed file version after a lost DB write, without rerendering."""
    for _ in range(3):
        detail = await repo.get_run(run_id)
        if detail is None or detail.status not in RUN_TERMINAL_STATUSES:
            return False
        if detail.orchestration is None:
            return True
        previous = detail.orchestration.checkpoint.get("scratch", {}).get(COMPLETION_KEY)
        if not isinstance(previous, dict):
            return True  # Historical records keep their original lifecycle contract.
        if detail.status not in {"done", "needs_review"}:
            return False
        version = await run_blocking(current_version, detail, settings.artifact_root)
        if version is None:
            return False
        if previous.get("content_version") == version:
            return True
        bundle = await run_rendering(load_version, detail, settings.artifact_root, version)
        record = await run_rendering(assess_completion, detail, bundle)
        if await repo.update_completion(
            run_id, record, expected_version=previous["content_version"]
        ):
            # A concurrent format retry may have published while we committed.
            if await run_blocking(current_version, detail, settings.artifact_root) == version:
                return True
    return False
