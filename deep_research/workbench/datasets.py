"""数据分析的结构化输入：CSV / TSV / XLSX → 一张明确来源的表。

约定：
- 表格在服务端解析，编码按 UTF-8（含 BOM）→ GB18030 依次尝试，Excel 导出的中文 CSV 也能读；
- 多工作表必须由用户明确选择，不替用户挑一张；
- 超过行、列或字符上限直接拒绝，不截断——截断后的表不能冒充完整的统计输入；
- 每张表都带来源（文件名、工作表、行列数、列类型），写进任务契约与报告。
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from .analysis import MAX_COLUMNS, MAX_ROWS, DatasetError, parse_dataset
from .contract import DATASET_MAX_CHARS

DATASET_MAX_FILE_BYTES = 16 * 1024 * 1024
TABLE_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx")


@dataclass(frozen=True)
class SheetTable:
    """一张可分析的表：``name`` 是工作表名，CSV / TSV 文件为空串。"""

    name: str
    csv: str
    rows: int
    columns: list[dict[str, str]] = field(default_factory=list)

    def profile(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "rows": self.rows,
            "columns": self.columns,
            "chars": len(self.csv),
        }


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DatasetError("无法识别文件编码，请另存为 UTF-8 的 CSV 后再上传")


def _column_type(series: Any) -> str:
    import pandas as pd

    if pd.api.types.is_bool_dtype(series):
        return "布尔"
    if pd.api.types.is_numeric_dtype(series):
        return "数值"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "日期"
    return "文本"


def profile_csv(name: str, text: str) -> SheetTable:
    """按分析时的同一解析规则统计行数与列类型；不合规时抛 ``DatasetError``。"""
    if len(text) > DATASET_MAX_CHARS:
        raise DatasetError(
            f"数据超过 {DATASET_MAX_CHARS:,} 字符上限，请先抽样或聚合后再上传（不会截断使用）"
        )
    frame = parse_dataset(text)
    columns = [{"name": str(column), "type": _column_type(frame[column])} for column in frame]
    return SheetTable(name=name, csv=text.strip(), rows=len(frame), columns=columns)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _xlsx_tables(raw: bytes) -> tuple[list[SheetTable], list[dict[str, str]]]:
    """逐个工作表转 CSV；空表跳过，不合规的表记下原因而不是丢弃整个文件。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # 可选依赖组 xlsx 未安装
        raise DatasetError("服务器未安装 Excel 解析组件，请改传 CSV / TSV") from exc

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对损坏文件抛出的异常类型不统一
        raise DatasetError("无法读取 Excel 文件，请确认是 .xlsx 格式且未加密") from exc
    tables: list[SheetTable] = []
    problems: list[dict[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            buffer = io.StringIO()
            writer = csv.writer(buffer, lineterminator="\n")
            count = 0
            too_large = False
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                count += 1
                if count > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                    too_large = True
                    break
                writer.writerow(cells)
            if too_large:
                problems.append(
                    {
                        "name": sheet.title,
                        "error": f"超过 {MAX_ROWS} 行或 {MAX_COLUMNS} 列上限，未截断使用",
                    }
                )
                continue
            if count == 0:
                continue
            try:
                tables.append(profile_csv(sheet.title, buffer.getvalue()))
            except DatasetError as exc:
                problems.append({"name": sheet.title, "error": str(exc)})
    finally:
        workbook.close()
    return tables, problems


def parse_table_file(raw: bytes, filename: str) -> dict[str, Any]:
    """解析上传的表格文件，返回每张可用工作表的 CSV 与概况。"""
    if not raw:
        raise DatasetError("文件为空")
    if len(raw) > DATASET_MAX_FILE_BYTES:
        raise DatasetError("文件超过 16 MB 限制")
    name = filename.casefold()
    if not name.endswith(TABLE_EXTENSIONS):
        raise DatasetError("仅支持 CSV、TSV 或 XLSX 表格")
    if name.endswith(".xlsx"):
        tables, problems = _xlsx_tables(raw)
    else:
        tables, problems = [profile_csv("", _decode(raw))], []
    if not tables:
        detail = "；".join(f"{item['name']}：{item['error']}" for item in problems)
        raise DatasetError("文件中没有可分析的表格" + (f"（{detail}）" if detail else ""))
    return {
        "filename": filename[:300],
        "sheets": [{**table.profile(), "csv": table.csv} for table in tables],
        "skipped": problems,
    }


def select_sheet(parsed: dict[str, Any], sheet: str | None) -> dict[str, Any]:
    """多工作表时必须指明用哪一张；不替用户挑选。"""
    sheets = parsed["sheets"]
    if sheet is None:
        if len(sheets) > 1:
            names = "、".join(item["name"] for item in sheets)
            raise DatasetError(f"文件包含多个工作表（{names}），请选择要分析的一张")
        return dict(sheets[0])
    for item in sheets:
        if item["name"] == sheet:
            return dict(item)
    raise DatasetError(f"找不到工作表「{sheet}」")


__all__ = [
    "DATASET_MAX_CHARS",
    "SheetTable",
    "parse_table_file",
    "profile_csv",
    "select_sheet",
]
