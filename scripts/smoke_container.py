"""Verify a running production image with only synthetic document contents.

Usage: docker compose exec -T api python - < scripts/smoke_container.py
"""

from __future__ import annotations

import io
import os
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

from deep_research.report.document import (
    ReportDocument,
    TableBlock,
    TableCell,
    TableColumn,
    TableRow,
)
from deep_research.report.pdf import render_pdf
from deep_research.report.xlsx import render_xlsx


def main() -> None:
    assert os.getuid() == 10001, "image must run as appuser"
    assert Path("/app/alembic/versions").is_dir()
    for resource, mime in (
        ("/readyz", "application/json"),
        ("/", "text/html"),
        ("/deep-research-icon.svg", "image/svg+xml"),
        ("/research-field.png", "image/png"),
    ):
        with urllib.request.urlopen("http://127.0.0.1:8000" + resource, timeout=5) as response:
            assert response.status == 200
            assert response.headers["Content-Type"].startswith(mime), resource
            assert response.read(), resource
    document = ReportDocument(
        query="生产镜像 XLSX / PDF 冒烟",
        blocks=[
            TableBlock(
                id="sample",
                columns=[TableColumn(key="value", label="Value")],
                rows=[TableRow(label="sample", cells={"value": TableCell(value="=1+1")})],
            )
        ],
    )
    workbook = load_workbook(io.BytesIO(render_xlsx(document)))
    assert workbook.active is not None
    values = [cell for row in workbook.active for cell in row if cell.value == "=1+1"]
    assert len(values) == 1 and values[0].data_type == "s"
    workbook.close()
    assert render_pdf(document).startswith(b"%PDF")
    print("Production image passed: non-root, live readiness/public assets, XLSX and Chinese PDF")


if __name__ == "__main__":
    main()
