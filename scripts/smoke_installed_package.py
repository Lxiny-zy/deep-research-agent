"""Smoke-test the installed wheel outside the source checkout."""

from __future__ import annotations

import asyncio
import io
import os
import tempfile
from pathlib import Path


def _resolved_path(value: str) -> Path:
    return Path(value).resolve()


def _path(value: str) -> Path:
    return Path(value)


async def _run() -> None:
    import httpx
    from openpyxl import load_workbook

    import deep_research
    from deep_research.api import _FRONTEND_ASSETS, _FRONTEND_DIST, app, lifespan
    from deep_research.intent.model import QUERY_MODEL_PATH
    from deep_research.persistence.db import _migration_root
    from deep_research.report.document import (
        ReportDocument,
        TableBlock,
        TableCell,
        TableColumn,
        TableRow,
    )
    from deep_research.report.xlsx import render_xlsx
    from deep_research.skills import SkillNotFoundError, default_skill_resolver

    source_root = _resolved_path(__file__).parents[1]
    package_path = _resolved_path(str(deep_research.__file__))
    if package_path == source_root / "deep_research" / "__init__.py":
        raise RuntimeError(
            f"smoke imported source checkout instead of installed wheel: {package_path}"
        )
    if not QUERY_MODEL_PATH.is_file():
        raise RuntimeError(f"intent model resource is missing: {QUERY_MODEL_PATH}")
    if not _FRONTEND_DIST.is_file() or not _FRONTEND_ASSETS.is_dir():
        raise RuntimeError(
            f"frontend resources are missing: index={_FRONTEND_DIST}, assets={_FRONTEND_ASSETS}"
        )
    migration_root = _migration_root()
    if not (migration_root / "alembic" / "versions").is_dir():
        raise RuntimeError(f"migration resources are missing: {migration_root}")
    source_migrations = {path.name for path in (source_root / "alembic" / "versions").glob("*.py")}
    installed_migrations = {
        path.name for path in (migration_root / "alembic" / "versions").glob("*.py")
    }
    missing_migrations = sorted(source_migrations - installed_migrations)
    if missing_migrations:
        raise RuntimeError(f"installed package migrations are incomplete: {missing_migrations}")

    with tempfile.TemporaryDirectory(prefix="deep-research-wheel-") as directory:
        root = _path(directory)
        isolated_project = root / "project"
        isolated_project.mkdir()
        resolver = default_skill_resolver(isolated_project)
        for skill_name in ("academic-search-v2", "pdf"):
            try:
                resolver.resolve(skill_name)
            except SkillNotFoundError as exc:
                raise RuntimeError(
                    f"installed package skill resource is missing: {skill_name}"
                ) from exc
        os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{root / 'smoke.db'}"
        os.environ["RUNTIME_CONFIG_PATH"] = str(root / "runtime_config.json")
        os.environ["DR_ARTIFACT_ROOT"] = str(root / "artifacts")
        os.environ["APP_ENV"] = "test"
        os.environ["API_KEY"] = ""
        os.environ["DR_API_KEYS"] = "[]"
        os.environ["DR_EXECUTION_MODE"] = "inline"
        async with lifespan(app):
            if not await app.state.repo.healthcheck():
                raise RuntimeError("installed package repository healthcheck failed")
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://wheel.test"
            ) as client:
                for resource, mime in (
                    ("/deep-research-icon.svg", "image/svg+xml"),
                    ("/research-field.png", "image/png"),
                ):
                    response = await client.get(resource)
                    assert response.status_code == 200
                    assert response.headers["content-type"].startswith(mime), resource
        if not (root / "smoke.db").is_file():
            raise RuntimeError("installed package did not initialize SQLite")
        document = ReportDocument(
            query="Wheel XLSX smoke",
            blocks=[
                TableBlock(
                    id="smoke",
                    columns=[TableColumn(key="value", label="Value")],
                    rows=[TableRow(label="sample", cells={"value": TableCell(value="=1+1")})],
                )
            ],
        )
        workbook = load_workbook(io.BytesIO(render_xlsx(document)))
        assert workbook.active is not None
        values = [cell for row in workbook.active for cell in row if cell.value == "=1+1"]
        assert len(values) == 1 and values[0].data_type == "s", "unsafe or incomplete XLSX export"
        workbook.close()


def main() -> None:
    asyncio.run(_run())
    print("installed package smoke passed")


if __name__ == "__main__":
    main()
