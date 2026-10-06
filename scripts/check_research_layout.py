"""Render current N1/N2 long-table fixtures without model calls or gate exemptions."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import time
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

from deep_research.bibliography import build_bibliography, present_markdown  # noqa: E402
from deep_research.models import (  # noqa: E402
    EvidenceVerification,
    ExperimentConditions,
    Finding,
    Quantity,
    Report,
    ResearchResult,
    Source,
)
from deep_research.persistence.repository import RunDetail  # noqa: E402
from deep_research.workbench.delivery.docx import render_docx  # noqa: E402
from deep_research.workbench.delivery.html import render_html  # noqa: E402
from deep_research.workbench.delivery.pdf import render_pdf  # noqa: E402
from deep_research.workbench.presentation_notes import append_presentation  # noqa: E402
from deep_research.workbench.research_presentation import research_context  # noqa: E402
from deep_research.workbench.support import evidence_id  # noqa: E402
from deep_research.workbench.tables import render_specs  # noqa: E402


def fixture(count: int = 36) -> tuple[RunDetail, dict]:
    findings, sources, rows = [], [], []
    for index in range(1, count + 1):
        label, value = f"M{index:03}", f"{0.8 + index / 1000:.3f}"
        statement = rf"合成方法采用误差表达式 $e_{{{index}}}=\frac{{a+b}}{{N}}$。"
        conditions = ExperimentConditions(
            dataset="Synthetic-A" if index % 2 else "Synthetic-B",
            split=f"test-split-{index:03}",
            train_data=f"training-group-{index:03}",
            bands=31 if index % 2 else 28,
            acquisition="模拟采集",
            spatial_size="128×128",
            protocol="训练裁剪尺寸；不代表完整测试输入尺寸",
        )
        quote = f"{label}: {statement} coverage >= {value} ± 0.001 ratio; {conditions.describe()}"
        source = Source(
            url=f"https://example.invalid/n2-fixture/paper-{index:03}",
            title=f"Reference-{index:03}. 合成长书目条目，用于验证来源目录跨页时仍保留完整文字、"
            "方法说明、版本与定位信息；不是外部论文，也不用于证明实际研究结论。"
            f" Synthetic comparison fixture volume {index}, version 2026-10-06.",
            content=quote,
            locator=f"synthetic page {index}, table row {index}",
        )
        finding = Finding(
            entity=label,
            statement=statement,
            evidence_quote=quote,
            source_url=source.url,
            quantity=Quantity(
                metric="coverage",
                value=float(value),
                rendered=value,
                comparator=">=",
                uncertainty=0.001,
                unit="ratio",
            ),
            conditions=conditions,
            verification=EvidenceVerification(
                status="verified",
                semantic_status="supported",
                quantity_status="verified",
                source_title=source.title,
                source_reference=source.title,
                source_content_hash=hashlib.sha256(quote.encode()).hexdigest(),
            ),
        )
        sources.append(source)
        findings.append(finding)
        rows.append(
            {
                "label": label,
                "cells": {
                    key: [evidence_id(finding)]
                    for key in (
                        "score",
                        "train",
                        "test",
                        "mechanism",
                    )
                },
            }
        )
    results = [ResearchResult(sub_question="合成比较表的训练、测试与误差表达式", findings=findings)]
    mapping = {source.url: index for index, source in enumerate(sources, 1)}
    spec = {
        "id": "n2-long-comparison",
        "title": "长比较表：条件与行内公式",
        "columns": [
            {"key": "score", "label": "比较值", "field": "quantity", "metric": "coverage"},
            {"key": "train", "label": "训练用途", "field": "conditions.train_data"},
            {"key": "test", "label": "测试用途", "field": "conditions.split"},
            {"key": "mechanism", "label": "方法表达式", "field": "statement"},
        ],
        "rows": rows,
    }
    source_markdown = (
        "## 工程样例说明\n\n"
        "本文件为受控合成排版材料，不代表论文结论，不证明研究覆盖或科学核验通过。"
        "按原始行序呈现，不以数值大小推断跨数据集排名。\n\n"
        "```evidence-table\n" + json.dumps(spec, ensure_ascii=False) + "\n```\n"
    )
    body, table_record = render_specs(source_markdown, results, mapping, comparison_context=True)
    detail = RunDetail(
        id="n2-layout-fixture",
        query="N2.4 确定性工程排版样例",
        status="needs_review",
        results=results,
        sources=sources,
        report=Report(query="N2.4", markdown=body, citations=list(mapping)),
        orchestration=SimpleNamespace(  # type: ignore[arg-type]
            checkpoint={"scratch": {"workbench": {"template": "litReview"}}},
        ),
    )
    return detail, table_record


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="docs/validation/n2-long-layout")
    args = parser.parse_args()
    destination = Path(args.output).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    started = time.monotonic()
    detail, table_record = fixture()
    assert detail.report
    original = detail.report.markdown
    display = append_presentation(original, detail)
    catalog = build_bibliography(
        display,
        detail.report.citations,
        detail.results[0].findings,
        detail.sources,
    )
    markdown = present_markdown(display, catalog, links=False)
    assert detail.report.markdown == original
    title = "N2 长表排版工程检查（合成样例）"
    version = hashlib.sha256(markdown.encode()).hexdigest()
    meta = "合成排版夹具 · 同一 Markdown 快照 · 无模型调用"
    assets = {
        "review.md": markdown.encode(),
        "review.html": render_html(markdown, title=title, meta=meta).encode(),
        "review.docx": render_docx(markdown, title=title, meta=meta),
        "review.pdf": render_pdf(markdown, title=title, meta=meta),
    }
    for name, data in assets.items():
        (destination / name).write_bytes(data)
    snapshot = {
        "fixture": "synthetic_only",
        "scientific_review": "not_performed",
        "input_markdown_sha256": version,
        "table_record": table_record,
        "research_context": research_context(detail),
        "sources": [source.model_dump(mode="json") for source in detail.sources],
    }
    (destination / "input-snapshot.json").write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    word = Document(io.BytesIO(assets["review.docx"]))
    word_xml = ET.fromstring(ZipFile(io.BytesIO(assets["review.docx"])).read("word/document.xml"))
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }
    word_text = "".join(element.text or "" for element in word_xml.findall(".//w:t", ns))
    equations = word_xml.findall(".//m:oMath", ns)
    assert len(equations) >= 36 and len(word.tables) == 2
    assert len(word.tables[0].rows) == 37 and len(word.tables[0].columns) == 5
    assert len(word_xml.findall(".//w:tblHeader", ns)) == 2
    assert not re.search(r"\\(?:frac|text|sum|begin)\b|\$e_", word_text)
    bounds_errors, table_pages, reference_pages, pages = [], [], [], []
    with pymupdf.open(stream=assets["review.pdf"], filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        squash = re.sub(r"\s+", "", text)
        for index in range(1, 37):
            value, label = f"{0.8 + index / 1000:.3f}", f"M{index:03}"
            assert value in word_text and value in squash
            assert label in word_text and label in squash
            assert f"Reference-{index:03}" in word_text and f"Reference-{index:03}" in squash
            assert word_text.count(value) == markdown.count(value) == squash.count(value)
            assert f"[{index}]" in word_text and f"[{index}]" in squash
        assert "训练裁剪尺寸；不代表完整测试输入尺寸" in squash
        assert "不证明整个领域不存在其他研究" in squash
        for number, page in enumerate(pdf, 1):
            page_text = page.get_text()
            compact = re.sub(r"\s+", "", page_text)
            if "方法表达式" in compact:
                assert all(
                    header in compact for header in ("对象", "比较值", "训练用途", "测试用途")
                )
                table_pages.append(number)
            if (
                "Reference-" in page_text
                and "参考文献" in text[: text.find(page_text) + len(page_text)]
            ):
                reference_pages.append(number)
            for block in page.get_text("dict")["blocks"]:
                if "lines" not in block:
                    continue
                for line in block["lines"]:
                    for span in line["spans"]:
                        rectangle = pymupdf.Rect(span["bbox"])
                        if (
                            rectangle.x0 < -0.5
                            or rectangle.y0 < -0.5
                            or rectangle.x1 > page.rect.width + 0.5
                            or rectangle.y1 > page.rect.height + 0.5
                        ):
                            bounds_errors.append(
                                {"page": number, "text": span["text"], "bbox": list(rectangle)}
                            )
            page.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(
                destination / f"page-{number:02}.png"
            )
            pages.append(
                {
                    "page": number,
                    "characters": len(page_text),
                    "width": page.rect.width,
                    "height": page.rect.height,
                    "images": len(page.get_image_info()),
                }
            )
        assert len(table_pages) >= 2, table_pages
        assert len(reference_pages) >= 2, reference_pages
        assert not bounds_errors, bounds_errors[:3]
    thumbnails = []
    for row in pages:
        with Image.open(destination / f"page-{row['page']:02}.png") as page_image:
            page_image.thumbnail((238, 337))
            thumbnail = Image.new("RGB", (258, 367), "#e2e8f0")
            thumbnail.paste(page_image, (10, 20))
            ImageDraw.Draw(thumbnail).text((10, 3), f"Page {row['page']}", fill="black")
            thumbnails.append(thumbnail)
    overview = Image.new("RGB", (258 * 4, 367 * ((len(thumbnails) + 3) // 4)), "white")
    for index, thumbnail in enumerate(thumbnails):
        overview.paste(thumbnail, (index % 4 * 258, index // 4 * 367))
    overview.save(destination / "pages-overview.png")
    summary = {
        "fixture": "synthetic_only",
        "model_calls": 0,
        "scientific_review": "not_performed",
        "input_markdown_sha256": version,
        "rows": 36,
        "comparison_columns": 5,
        "word_tables": len(word.tables),
        "word_equations": len(equations),
        "word_repeat_headers": 2,
        "pdf_pages": len(pages),
        "pdf_table_header_pages": table_pages,
        "pdf_reference_pages": reference_pages,
        "pdf_out_of_page_text": bounds_errors,
        "numeric_occurrences_match_markdown_word_pdf": True,
        "pdf_math_note": "Vector math uses source TeX as accessibility ActualText; "
        "screenshots verify visible typesetting.",
        "assets": [
            {"name": name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
            for name, data in assets.items()
        ],
        "pages": pages,
        "elapsed_seconds": round(time.monotonic() - started, 2),
    }
    (destination / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {key: value for key, value in summary.items() if key not in {"assets", "pages"}},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
