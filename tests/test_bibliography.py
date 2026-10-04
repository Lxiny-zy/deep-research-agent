from __future__ import annotations

import hashlib

from deep_research.bibliography import build_bibliography, document_identity, present_markdown
from deep_research.models import Source, SourceIdentity
from tests.fakes import verified_finding


def test_documents_are_grouped_but_every_location_and_hash_survives():
    urls = [f"https://workspace.invalid/attachments/paper-a?chunk={i}" for i in (1, 2, 3)]
    other = "https://example.org/article?id=2"
    findings = [verified_finding(source_url=url) for url in [*urls, other]]
    original = "B [4]。A 方法 [2]，A 结果 [1,3]。\n\n## 参考来源\n[1] old"
    catalog = build_bibliography(original, [*urls, other], findings)
    assert len(catalog.documents) == 2
    assert catalog.documents[0].url == other  # order of first use, not source-list order
    assert catalog.documents[1].locations == [1, 2, 3]
    assert [(p.index, p.document, p.url) for p in catalog.locations] == [
        (1, 2, urls[0]),
        (2, 2, urls[1]),
        (3, 2, urls[2]),
        (4, 1, other),
    ]
    assert all(p.content_hashes == ["fixture-hash"] for p in catalog.locations)
    assert "A 方法 [[2]](#cite-2)" in catalog.body
    assert "A 结果 [[2]](#cite-1-3)" in catalog.body
    assert "B [[1]](#cite-4)" in catalog.body
    projected = present_markdown(original, catalog)
    assert projected.count("## 参考文献") == 1 and "old" not in projected
    assert present_markdown(projected, catalog) == projected


def test_numeric_interval_is_not_counted_or_rewritten_as_a_bibliographic_location():
    from deep_research.workbench.delivery.math_markdown import citation_text
    from deep_research.workbench.prose_review import _CITE, prose_units

    text = "RGB images are rescaled to [0,1] [2]. 方法参见 [1,2]。"
    urls = ["https://paper.test/a", "https://paper.test/b"]
    catalog = build_bibliography(text, urls, [])
    assert "[0,1]" in catalog.body and "#cite-0" not in catalog.body
    assert _CITE.findall(citation_text(text)) == ["2", "1,2"]
    units, _ = prose_units(text, [1, 2])
    assert units[0].citations == [1, 2]


def test_same_filename_is_not_document_identity_and_versions_are_preserved():
    urls = [
        "https://workspace.invalid/attachments/first?chunk=1",
        "https://workspace.invalid/attachments/second?chunk=1",
    ]
    findings = [verified_finding(source_url=u) for u in urls]
    for finding in findings:
        finding.verification.source_title = "paper.pdf"
    assert len(build_bibliography("A [1,2]", urls, findings).documents) == 2
    assert (
        document_identity("https://arxiv.org/pdf/2401.01234v1.pdf")[0]
        == document_identity("https://arxiv.org/html/2401.01234v1?dr_section=2")[0]
    )
    assert (
        document_identity("https://arxiv.org/abs/2401.01234v2")[0]
        != document_identity("https://arxiv.org/abs/2401.01234v1")[0]
    )


def test_reparsed_immutable_pdf_shares_bibliography_but_keeps_both_evidence_snapshots():
    base = "https://workspace.invalid/attachments/fe8aaa9708c6562af5647757"
    old = base + "?chunk=16"
    new = old + "&text_revision=0123456789abcdef"
    findings = [verified_finding(source_url=url) for url in (old, new)]
    findings[0].verification.source_content_hash = "old-text"
    findings[1].verification.source_content_hash = "superscript-text"
    catalog = build_bibliography("Earlier finding [1]. Correct number [2].", [old, new], findings)
    assert len(catalog.documents) == 1
    assert [(p.url, p.content_hashes) for p in catalog.locations] == [
        (old, ["old-text"]),
        (new, ["superscript-text"]),
    ]
    assert "Correct number [[1]](#cite-2)" in catalog.body
    for variant in (
        old + "&version=2",
        old + "&text_revision=unknown",
        new.replace("workspace.invalid", "example.org"),
        new.replace("/attachments/", "/sources/"),
    ):
        assert document_identity(variant)[0] != document_identity(base)[0]


def test_article_selectors_and_unknown_fragments_are_not_dropped():
    a = "https://example.org/article?id=1&dr_section=pdf-1"
    b = "https://example.org/article?id=1&dr_section=pdf-8"
    c = "https://example.org/article?id=2&dr_section=pdf-1"
    assert document_identity(a)[0] == document_identity(b)[0]
    assert document_identity(a)[0] != document_identity(c)[0]
    assert (
        document_identity("https://example.org/app#/paper/1")[0]
        != document_identity("https://example.org/app#/paper/2")[0]
    )
    assert (
        document_identity("https://example.org/p#chunk-1")[0]
        == document_identity("https://example.org/p#chunk-2")[0]
    )


def test_projection_preserves_math_code_links_and_invalid_citations():
    urls = ["https://example.org/a", "https://example.org/a#chunk-2"]
    text = "A [2]。$X[2]$ `keep [1]` [link](https://example.org/[2])。unknown [99]"
    cat = build_bibliography(text, urls, [])
    assert "A [[1]](#cite-2)" in cat.body
    assert "$X[2]$ `keep [1]` [link](https://example.org/[2])。unknown [99]" in cat.body
    assert cat.source_body == text
    adjacent = build_bibliography("A [1][2]. B [1] [2]. C [1], [2].", urls, [])
    assert adjacent.body == "A [[1]](#cite-1-2). B [[1]](#cite-1-2). C [[1]](#cite-1-2)."


def test_duplicate_query_parameters_keep_their_document_selecting_order():
    assert (
        document_identity("https://example.org/article?id=1&id=2")[0]
        != document_identity("https://example.org/article?id=2&id=1")[0]
    )


def test_explicit_publisher_citation_is_taken_only_from_matching_snapshot():
    url = "https://workspace.invalid/attachments/abc?chunk=1"
    content = (
        "Citation: Author, A. A Study.\nJournal 2024, 12. https://doi.org/\n"
        "10.1234/study\nAcademic Editor: Editor\nOther doi 10.1234/other"
    )
    source = Source(url=url, content=content, locator="A long parser heading / 第 1-2 页 / 片段 1")
    finding = verified_finding(source_url=url)
    finding.verification.source_content_hash = hashlib.sha256(content.encode()).hexdigest()
    cat = build_bibliography("A [1]", [url], [finding], [source])
    assert (
        cat.documents[0].reference
        == "Author, A. A Study. Journal 2024, 12. https://doi.org/10.1234/study"
    )
    assert cat.documents[0].url == "https://doi.org/10.1234/study"
    assert cat.locations[0].label == "第 1-2 页 · 片段 1"
    changed = source.model_copy(update={"content": content + "changed"})
    other = build_bibliography("A [1]", [url], [finding], [changed])
    assert "2024" not in other.documents[0].reference
    later = source.model_copy(update={"locator": "第 9 页 / 引用格式示例"})
    assert (
        "2024" not in build_bibliography("A [1]", [url], [finding], [later]).documents[0].reference
    )


def test_external_doi_aliases_group_without_using_title_similarity():
    urls = ["https://publisher.example/article/1?dr_section=pdf-2", "https://doi.org/10.1234/paper"]
    findings = [verified_finding(source_url=url) for url in urls]
    findings[0].verification.source_identity = SourceIdentity(doi="10.1234/paper")
    cat = build_bibliography("A [1,2]", urls, findings)
    assert len(cat.documents) == 1 and "(#cite-1-2)" in cat.body
    from deep_research.bibliography import work_keys
    from deep_research.workbench.scholarly import source_counts

    assert source_counts(urls, document_keys=work_keys(cat)) == (1, [])
    local = [
        "https://workspace.invalid/attachments/one?chunk=1",
        "https://workspace.invalid/attachments/two?chunk=1",
    ]
    assert source_counts(local, {url: "same-file.pdf" for url in local}) == (2, [])


def test_presentation_does_not_rewrite_checked_document_or_table_values():
    from deep_research.models import Report, ResearchResult
    from deep_research.report.assemble import assemble_document
    from deep_research.report.document import TableBlock, TableCell, TableColumn, TableRow
    from deep_research.report.presentation import presentation_document

    urls = [f"https://workspace.invalid/attachments/p?chunk={i}" for i in (1, 2)]
    doc = assemble_document(
        Report(query="q", markdown="Result [2]", citations=urls),
        [ResearchResult(sub_question="q", findings=[verified_finding(source_url=u) for u in urls])],
        extra_blocks=[
            TableBlock(
                id="data",
                columns=[TableColumn(key="value", label="Value")],
                rows=[
                    TableRow(
                        label="Alpha",
                        citation=2,
                        cells={
                            "value": TableCell(
                                value="38.4", numeric=38.4, citations=[1, 2], note_ref=7
                            )
                        },
                    )
                ],
            )
        ],
    )
    before = doc.model_dump_json()
    view = presentation_document(doc)
    assert doc.model_dump_json() == before
    assert len(view.references) == 1 and view.blocks[0].markdown == "Result [1]"
    table = view.table("data")
    assert table.rows[0].citation == 1
    assert table.rows[0].cells["value"].citations == [1]
    assert table.rows[0].cells["value"].value == "38.4"
    assert table.rows[0].cells["value"].note_ref == 7
    assert [r.source_url for r in view.evidence] == [r.source_url for r in doc.evidence]


def test_offline_html_links_reveal_only_the_requested_locations_without_scripts():
    import re

    from deep_research.workbench.delivery.html import render_html

    urls = [f"https://workspace.invalid/attachments/p?chunk={i}" for i in (1, 2, 3)]
    catalog = build_bibliography("One [2]. Combined [1,3].", urls, [])
    evidence = [
        {"citation": i, "statement": f"claim {i}", "quote": f"QUOTE{i} <script>bad</script>"}
        for i in (1, 2, 3)
    ]
    html = render_html(
        present_markdown(catalog.source_body, catalog),
        title="Report",
        bibliography=catalog,
        evidence=evidence,
    )
    combined = re.search(
        r'<aside class="citation-location" id="cite-source-1-3">(.*?)</aside>', html, re.S
    )[1]
    assert "QUOTE1" in combined and "QUOTE3" in combined and "QUOTE2" not in combined
    assert 'href="#cite-2"' in html and 'id="cite-2"' in html
    assert "<script" not in html and "&lt;script&gt;" in html


def test_shared_table_cells_are_projected_once_and_view_is_idempotent():
    from deep_research.report.document import ReportDocument, TableBlock, TableCell, TableRow
    from deep_research.report.presentation import presentation_document

    urls = [
        "https://example.org/b",
        "https://example.org/a#chunk-1",
        "https://example.org/a#chunk-2",
    ]
    catalog = build_bibliography("A [2]. B [1].", urls, [])
    cell = TableCell(value="value", citations=[2])
    table = TableBlock(
        id="t",
        rows=[TableRow(label="one", cells={"v": cell}), TableRow(label="two", cells={"v": cell})],
    )
    document = ReportDocument(blocks=[table], bibliography=catalog)
    view = presentation_document(document)
    assert [row.cells["v"].citations for row in view.table("t").rows] == [[1], [1]]
    assert cell.citations == [2]
    assert presentation_document(view) is view


def test_reference_slide_pagination_keeps_long_entries_and_every_character():
    import io
    import re

    from pptx import Presentation

    from deep_research.workbench.delivery.pptx import reference_pages, render_pptx

    long_entry = "完整参考文献内容" * 500
    pages = reference_pages([long_entry])
    assert len(pages) > 1
    assert (
        "".join(re.sub(r"^\[1\] (?:（续）)?", "", text) for page in pages for text in page)
        == long_entry
    )
    entries = [f"Author-{i}. " + "完整论文题名和出版信息" * 50 for i in range(10)]
    data = render_pptx({"title": "Bibliography", "slides": []}, citations=entries)
    deck = Presentation(io.BytesIO(data))
    reference_slides = [
        slide
        for slide in deck.slides
        if any(shape.has_text_frame and shape.text.startswith("参考文献") for shape in slide.shapes)
    ]
    assert len(reference_slides) > 1
    text = "\n".join(
        shape.text for slide in reference_slides for shape in slide.shapes if shape.has_text_frame
    )
    assert all(f"Author-{i}." in text for i in range(10))
    assert "[10]" in text


def test_mindmap_displays_document_numbers_without_merging_location_evidence():
    import re

    from deep_research.workbench.delivery.mindmap import render_mindmap_html, render_svg

    urls = [f"https://workspace.invalid/attachments/p?chunk={i}" for i in (1, 2, 3)]
    catalog = build_bibliography("node [1,3]", urls, [])
    value = {
        "root": "Topic",
        "branches": [
            {"label": "Claim", "citations": [1, 3], "display_citations": [1], "children": []}
        ],
        "sources": [
            {"index": i, "url": url, "quotes": [f"QUOTE{i}"]} for i, url in enumerate(urls, 1)
        ],
        "bibliography": catalog.model_dump(mode="json"),
        "evidence": [
            {"citation": i, "statement": "Claim", "quote": f"QUOTE{i}"} for i in (1, 2, 3)
        ],
    }
    assert "Claim[1]" in render_svg(value) and "Claim[1][3]" not in render_svg(value)
    html = render_mindmap_html(value, title="Map")
    linked = re.search(r'<aside class="citation-location" id="cite-1-3">(.*?)</aside>', html, re.S)[
        1
    ]
    assert "未绑定依据" in linked and "QUOTE1" not in linked
    linked = re.search(
        r'<aside class="citation-location" id="cite-source-1-3">(.*?)</aside>', html, re.S
    )[1]
    assert "QUOTE1" in linked and "QUOTE3" in linked and "QUOTE2" not in linked


def test_reference_section_removal_preserves_appendix_and_checks_its_numbers():
    from deep_research.bibliography import source_body
    from deep_research.models import ResearchResult
    from deep_research.report.validation import validate_body

    text = (
        "A [1].\n\n```text\n## References\nkeep code\n```\n\n"
        "## 参考来源\n[1] url\n\n## 附录\nMissing 999 [1]."
    )
    body = source_body(text)
    assert "## References\nkeep code" in body and "[1] url" not in body
    assert "## 附录\nMissing 999 [1]." in body
    finding = verified_finding(source_url="https://example.org/paper")
    check = validate_body(
        text,
        [ResearchResult(sub_question="q", findings=[finding])],
        {finding.source_url: 1},
        fallback=False,
    )
    assert any("999" in str(problem) for problem in check.problems)
    assert (
        source_body("A\r\n\r\n## References\r\nref\r\n\r\n## Appendix\r\nB")
        == "A\r\n\r\n## Appendix\r\nB"
    )


async def test_old_prose_review_cannot_authorize_an_appended_section_after_references():
    from deep_research.models import ResearchResult
    from deep_research.workbench.prose_review import ProseReviewer
    from tests.fakes import FakeLLM

    finding = verified_finding(source_url="https://example.org/paper")
    llm = FakeLLM()
    reviewer = ProseReviewer.research(
        llm, [ResearchResult(sub_question="q", findings=[finding])], {finding.source_url: 1}, 50000
    )
    original = "Original claim [1].\n\n## 参考来源\n[1] paper"
    record = await reviewer.review(original)
    calls = llm.parse_calls
    bound, issues = reviewer.check(original + "\n\n## 附录\nAdditional claim [1].", record)
    assert not bound and issues
    assert llm.parse_calls == calls


def test_line_broken_doi_is_not_silently_published_as_a_different_identifier():
    url = "https://workspace.invalid/attachments/paper?chunk=1"
    source = Source(
        url=url,
        content="Citation: Author. Study. https://doi.org/10.1234/long\nsuffix\nReceived: today",
    )
    finding = verified_finding(source_url=url)
    finding.verification.source_content_hash = hashlib.sha256(source.content.encode()).hexdigest()
    catalog = build_bibliography("A [1]", [url], [finding], [source])
    assert "10.1234/long" not in catalog.documents[0].reference
