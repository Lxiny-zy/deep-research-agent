from deep_research.workbench.analysis import analyse, check_numbers, fallback_report


def test_declared_input_row_numbers_do_not_authorize_statistical_values():
    question = "第 12 条记录的测量值人为置空。"
    facts = "均值=5.8, n=149"
    assert (
        check_numbers("输入说明：第12行人为置空。均值为5.8。", facts, input_description=question)
        == []
    )
    assert check_numbers("第12条记录已置空，均值为12。", facts, input_description=question) == [
        "12"
    ]
    assert check_numbers("第99行人为置空。", facts, input_description=question) == ["99"]


def test_fallback_reports_only_executed_group_methods_and_declared_correction_policy():
    data = "group,value\nA,1\nA,2\nA,3\nB,4\nB,5\nB,6\nC,7\nC,8\nC,9\n"
    result = analyse(data, "比较三组")
    body = fallback_report(result)
    assert "单因素方差分析" in body and "Welch" not in body
    assert "Kruskal" in body and "多重比较校正" in result.facts()
    assert "未执行" in result.facts()
    assert float(result.tests[0]["eta_squared"]) == 0.9
    assert [group["n"] for group in result.tests[0]["group_summaries"]] == [3, 3, 3]
    assert result.tests[0]["n_total"] == 9 and result.tests[0]["df_within"] == 6


def test_group_counts_follow_variable_specific_missingness_and_are_exported():
    import io

    from openpyxl import load_workbook

    from deep_research.workbench.publish import _stats_xlsx

    result = analyse(
        "group,x,y\nA,1,2\nA,2,3\nA,,4\nB,4,5\nB,5,6\nB,6,7\nC,7,8\nC,8,9\nC,9,10\n", "比较三组"
    )
    x = next(test for test in result.tests if test["variable"] == "x")
    y = next(test for test in result.tests if test["variable"] == "y")
    assert [g["n"] for g in x["group_summaries"]] == [2, 3, 3]
    assert [g["n"] for g in y["group_summaries"]] == [3, 3, 3]
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    assert "分组统计" in workbook.sheetnames
    assert "eta_squared" in [cell.value for cell in workbook["显著性检验"][1]]


def test_unavailable_test_is_not_reported_as_no_significant_difference():
    result = analyse("group,value\nA,1\nA,2\nB,3\nB,4\n", "比较两组")
    result.tests[0].update(
        significant=None, statistic=None, p_value=None, robust_method=None, robust_p=None
    )
    body = fallback_report(result)
    assert "未完成差异检验" in body and "未发现显著差异" not in body
    assert "None p=None" not in body


def test_paired_and_independent_tests_keep_their_own_assumption_statements():
    paired = analyse("a,b\n1,2\n2,4\n3,4\n4,7\n", "配对比较")
    independent = analyse("group,value\nA,1\nA,2\nA,3\nB,3\nB,5\nB,7\n", "比较两组")
    assert "各配对对象独立" in paired.facts()
    assert "各配对对象独立" not in independent.facts()
    assert "不要求两组方差相等" in independent.facts()
    assert "方差分析自由度" not in independent.facts()


def test_correlation_counts_pairwise_complete_rows_in_ledger_and_export():
    import io

    from openpyxl import load_workbook

    from deep_research.workbench.publish import _stats_xlsx

    result = analyse("x,y,z\n1,2,3\n2,4,5\n3,,6\n4,7,\n5,8,9\n6,9,10\n")
    counts = {(r["a"], r["b"]): r["n"] for r in result.correlations}
    assert counts == {("x", "y"): 5, ("x", "z"): 5, ("y", "z"): 4}
    assert "有效样本量 n=4" in result.facts()
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    rows = list(workbook["相关性"].values)
    assert rows[0] == ("a", "b", "r", "p_value", "n")
    assert {(r[0], r[1]): r[-1] for r in rows[1:]} == counts


def test_legacy_correlation_snapshot_does_not_invent_unrecorded_counts():
    from deep_research.workbench.analysis import ledger_facts

    result = analyse("x,y\n1,2\n2,4\n3,5\n4,7\n")
    snapshot = result.snapshot()
    snapshot.pop("facts")
    for row in snapshot["correlations"]:
        row.pop("n")
    assert "有效样本量 n=" not in ledger_facts(snapshot)
