"""Inferential methods, familywise comparisons and scoped interpretation."""

from pathlib import Path

import pytest

from deep_research.workbench.analysis import analyse
from deep_research.workbench.analysis_scope import AnalysisScope


@pytest.fixture
def penguins():
    csv = (Path(__file__).parent / "fixtures" / "penguins_sc75.csv").read_text(encoding="utf-8")
    return analyse(
        csv,
        "比较物种的鳍长和体重",
        scope=AnalysisScope(
            measures=["flipper_length_mm", "body_mass_g"],
            groups=["species"],
            background=["sex", "year", "island"],
        ).model_dump(),
    )


def test_heterogeneous_penguin_mass_uses_welch_and_games_howell(penguins):
    test = next(row for row in penguins.tests if row["variable"] == "body_mass_g")
    assert test["method"] == "Welch ANOVA"
    variance = next(row for row in test["assumptions"] if row["kind"] == "variance_homogeneity")
    assert variance["p_value"] == pytest.approx(0.00644508, abs=0.00001)
    assert variance["status"] == "rejected"
    assert len([row for row in test["assumptions"] if row["kind"] == "normality"]) == 3
    pair = next(
        row
        for row in test["posthoc"]
        if {row["left_group"], row["right_group"]} == {"Adelie", "Chinstrap"}
    )
    assert pair["method"] == "Games-Howell"
    assert pair["p_value"] == pytest.approx(0.8501540, abs=0.0001)
    assert pair["significant"] is False
    assert pair["ci_low"] < 0 < pair["ci_high"]


def test_equal_variances_use_anova_and_tukey_with_declared_difference_direction():
    csv = "group,value\n" + "\n".join(
        f"{label},{offset + value}"
        for label, offset in [("A", 0), ("B", 2), ("C", 10)]
        for value in [1, 2, 3, 4, 5]
    )
    result = analyse(csv)
    test = result.tests[0]
    assert test["method"] == "单因素方差分析"
    assert len(test["posthoc"]) == 3
    assert all(row["method"] == "Tukey HSD" for row in test["posthoc"])
    pair = next(
        row for row in test["posthoc"] if row["left_group"] == "A" and row["right_group"] == "C"
    )
    assert pair["mean_difference"] == 10
    assert pair["ci_low"] > 0 and pair["ci_high"] > pair["ci_low"]


def test_within_species_correlations_are_not_replaced_by_the_pooled_value(penguins):
    within = {
        row["group"]: row for row in penguins.correlations if row.get("group_column") == "species"
    }
    assert set(within) == {"Adelie", "Chinstrap", "Gentoo"}
    for name, expected in {
        "Adelie": 0.46820169,
        "Chinstrap": 0.64155941,
        "Gentoo": 0.70266652,
    }.items():
        assert within[name]["r"] == pytest.approx(expected, abs=0.0001)
    assert [within[name]["n"] for name in ["Adelie", "Chinstrap", "Gentoo"]] == [151, 68, 123]


async def test_omnibus_significance_cannot_claim_all_pairs_differ(penguins):
    from deep_research.workbench.prose_review import reviewer_for_report
    from tests.fakes import FakeLLM

    checker = reviewer_for_report(
        FakeLLM(),
        "比较物种体重",
        [],
        [],
        {
            "workbench": {"template": "dataAnalysis"},
            "analysis": penguins.snapshot(),
        },
        200000,
    )
    audit = await checker.review("三个物种的 body_mass_g 都不同。")
    assert audit["status"] == "fail"
    assert any("Adelie" in issue and "Chinstrap" in issue for issue in audit["issues"])


def test_pairwise_claims_respect_adjusted_significance_and_direction(penguins):
    from deep_research.workbench.analysis_review import posthoc_scope_issues

    ledger = penguins.snapshot()
    assert posthoc_scope_issues("Adelie 的体重显著高于 Gentoo。", ledger)
    assert not posthoc_scope_issues("Gentoo 的体重显著高于 Adelie 和 Chinstrap。", ledger)
    assert not posthoc_scope_issues("三个物种的体重存在总体差异。", ledger)
    assert not posthoc_scope_issues("Adelie 与 Chinstrap 的体重差异未达统计显著。", ledger)
    assert not posthoc_scope_issues("Adelie 与 Chinstrap 的体重没有显著差异。", ledger)
    assert posthoc_scope_issues("三个物种的 body_mass_g 均存在显著差异。", ledger)
    assert posthoc_scope_issues("All species differ in body_mass_g.", ledger)


def test_nonsignificant_difference_does_not_prove_equality(penguins):
    from deep_research.workbench.analysis_review import posthoc_scope_issues

    assert posthoc_scope_issues("Adelie 与 Chinstrap 的体重相同。", penguins.snapshot())
    assert posthoc_scope_issues("Adelie 与 Chinstrap 的体重没有差异。", penguins.snapshot())
    assert not posthoc_scope_issues(
        "Adelie 与 Chinstrap 的体重未达显著差异，但这不证明两者相同。", penguins.snapshot()
    )
    assert posthoc_scope_issues("Adelie 与 Chinstrap 的体重样本均值相同。", penguins.snapshot())


def test_sample_mean_equality_needs_recorded_means():
    from deep_research.workbench.analysis_review import posthoc_scope_issues

    ledger = {"tests": [{"variable": "mass", "group": "group", "group_summaries": [
        {"label": "A", "n": 3, "mean": 0}, {"label": "B", "n": 3, "mean": 0},
    ]}]}
    assert not posthoc_scope_issues("A 与 B 组的 mass 样本均值相同。", ledger)
    for row in ledger["tests"][0]["group_summaries"]:
        row.pop("mean")
    assert posthoc_scope_issues("A 与 B 组的 mass 样本均值相同。", ledger)


def test_posthoc_failure_does_not_discard_the_completed_omnibus_test(monkeypatch):
    from deep_research.workbench import analysis_methods
    from deep_research.workbench.analysis_review import posthoc_scope_issues

    def broken(*args, **kwargs):
        raise ValueError("cannot estimate pairwise interval")

    monkeypatch.setattr(analysis_methods, "posthoc_comparisons", broken)
    result = analyse("group,value\nA,1\nA,2\nA,3\nB,5\nB,6\nB,7\nC,9\nC,10\nC,11")
    assert result.tests[0]["p_value"] != "NA"
    assert result.tests[0]["posthoc"] == []
    assert result.tests[0]["posthoc_error"] and result.issues
    assert posthoc_scope_issues("所有组的 value 都不同。", result.snapshot())


def test_penguin_facts_and_workbook_keep_diagnostics_posthoc_and_group_correlation_scopes(penguins):
    import io

    from openpyxl import load_workbook

    from deep_research.workbench.analysis import fallback_report
    from deep_research.workbench.analysis_review import statistic_scope_issues
    from deep_research.workbench.publish import _stats_xlsx

    assert statistic_scope_issues(penguins.facts(), penguins.snapshot()) == []
    assert statistic_scope_issues(fallback_report(penguins), penguins.snapshot()) == []
    book = load_workbook(io.BytesIO(_stats_xlsx(penguins)))
    assert {"统计前提", "事后比较", "相关性"} <= set(book.sheetnames)
    comparisons = list(book["事后比较"].values)
    assert len(comparisons) == 7
    assert "confidence_level" in comparisons[0] and "adjustment" in comparisons[0]
    correlations = list(book["相关性"].values)
    assert len(correlations) == 5
    assert [row[correlations[0].index("group")] for row in correlations[1:]] == [
        None,
        "Adelie",
        "Chinstrap",
        "Gentoo",
    ]


def test_frozen_statistics_do_not_recompute_or_add_methods(penguins, monkeypatch):
    import copy

    from deep_research.workbench import analysis_methods

    def forbidden(*args, **kwargs):
        raise AssertionError("frozen statistics must not be computed again")

    for name in (
        "group_assumptions",
        "normality",
        "posthoc_comparisons",
        "within_group_correlations",
    ):
        monkeypatch.setattr(analysis_methods, name, forbidden)
    csv = (Path(__file__).parent / "fixtures" / "penguins_sc75.csv").read_text(encoding="utf-8")
    frozen = penguins.snapshot()
    restored = analyse(csv, frozen=frozen)
    assert restored.snapshot() == frozen and restored.figures == penguins.figures
    legacy = copy.deepcopy(frozen)
    for row in legacy["tests"]:
        row.pop("assumptions", None)
        row.pop("posthoc", None)
        row.pop("posthoc_error", None)
    legacy["correlations"] = [row for row in legacy["correlations"] if "group" not in row]
    restored = analyse(csv, frozen=legacy)
    assert restored.tests == legacy["tests"] and restored.correlations == legacy["correlations"]
    assert all("posthoc" not in row for row in restored.tests)


def test_small_or_constant_samples_do_not_confirm_normality():
    from deep_research.workbench.analysis_methods import normality

    for values in ([1, 2], [1, 1, 1, 1]):
        result = normality(values, variable="value")
        assert result["status"] == "unavailable" and result["p_value"] == "NA"


def test_unavailable_diagnostic_note_does_not_claim_the_assumption_holds():
    from deep_research.workbench.analysis import fallback_report
    from deep_research.workbench.analysis_review import statistic_scope_issues

    result = analyse("group,value\nA,1\nA,2\nB,3\nB,4")
    assert statistic_scope_issues(result.facts(), result.snapshot()) == []
    assert statistic_scope_issues(fallback_report(result), result.snapshot()) == []


@pytest.mark.parametrize("value", [0.04999999999, 0.05000000001, 0.00099999999999])
def test_displayed_p_value_does_not_cross_the_decision_threshold(value):
    from deep_research.workbench.analysis import _fmt_p

    displayed = _fmt_p(value)
    for threshold in (0.05, 0.01, 0.001):
        assert (displayed < threshold) == (value < threshold)


def test_rejected_diagnostics_cannot_be_reported_as_proven_assumptions(penguins):
    from deep_research.workbench.analysis_review import assumption_scope_issues

    ledger = penguins.snapshot()
    assert assumption_scope_issues("body_mass_g 的方差齐性已经成立。", ledger)
    assert assumption_scope_issues("body_mass_g 各组均通过正态性检验。", ledger)
    assert assumption_scope_issues("Shapiro-Wilk 检验通过，证明观测独立。", ledger)
    assert not assumption_scope_issues("未拒绝正态性不等于证明前提成立。", ledger)
