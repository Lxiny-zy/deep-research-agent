"""The workbook keeps complete input counts and distinguishes file and analysis hashes."""

import base64
import hashlib
import io
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from deep_research.workbench.analysis import analyse
from deep_research.workbench.publish import _stats_xlsx
from tests.test_datasets import _client, _contract, _xlsx
from tests.test_datasets import api_repo as api_repo


def _digest(text):
    return hashlib.sha256(text.strip().encode("utf-8")).hexdigest()


def _overview(workbook):
    return {row[0]: row[1] for row in list(workbook["输入概况"].values)[1:]}


def test_workbook_counts_every_column_even_when_not_selected_for_analysis():
    from tests.test_analysis_scope import CSV, SCOPE

    result = analyse(CSV, scope=SCOPE.model_dump())
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    overview = _overview(workbook)
    assert overview["总行数"] == 6 and overview["总列数"] == 7
    assert overview["分析输入文本 SHA-256"] == _digest(CSV)
    assert overview["原始文件字节 SHA-256"] == "未记录"
    assert "不代表独立样本量" in overview["行数口径"]
    counts = {row[0]: row[1:] for row in list(workbook["缺失概况"].values)[1:]}
    assert counts == {
        "id": (6, 6, 0), "species": (6, 6, 0), "year": (6, 6, 0),
        "sex": (6, 5, 1), "mass": (6, 5, 1), "length": (6, 5, 1), "unused": (6, 6, 0),
    }
    assert "分析提示" in workbook.sheetnames and "未完成分析" not in workbook.sheetnames


def test_legacy_workbook_marks_unknown_missing_counts_and_hashes_as_unrecorded():
    result = SimpleNamespace(
        rows=2, columns=["x", "y"], missing={"x": 0},
        describe=[], tests=[], correlations=[], issues=[],
    )
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    assert list(workbook["缺失概况"].values)[1:] == [
        ("x", 2, 2, 0), ("y", 2, "未记录", "未记录"),
    ]
    assert _overview(workbook)["分析输入文本 SHA-256"] == "未记录"
    assert _overview(workbook)["原始文件字节 SHA-256"] == "未记录"


def test_frozen_restore_does_not_invent_missing_metadata_from_new_parser_rules():
    csv = "x,y\n1,2\n2,3\n3,4"
    frozen = analyse(csv).snapshot()
    frozen["missing"].pop("y")
    frozen.pop("input_sha256")
    restored = analyse(csv, frozen=frozen)
    workbook = load_workbook(io.BytesIO(_stats_xlsx(restored)))
    assert list(workbook["缺失概况"].values)[2] == ("y", 3, "未记录", "未记录")
    assert _overview(workbook)["分析输入文本 SHA-256"] == "未记录"


def test_workbook_saves_names_and_notices_as_text_instead_of_excel_formulas():
    result = analyse("x\n1\n2\n3", source={"filename": "=file.xlsx", "sheet": "=Sheet"})
    result.columns = ["=x"]
    result.describe[0]["variable"] = "=x"
    result.missing = {"=x": 0}
    result.issues = ["=notice"]
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)), data_only=False)
    assert workbook["描述统计"]["A2"].value == "=x"
    assert workbook["描述统计"]["A2"].data_type == "s"
    assert _overview(workbook)["来源文件"] == "=file.xlsx"
    assert _overview(workbook)["工作表"] == "=Sheet"
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


@pytest.mark.parametrize("filename,raw", [
    ("中文.csv", "列\r\n1\r\n2\r\n".encode("gb18030")),
    ("bom.csv", b"\xef\xbb\xbfx\r\n1\r\n2\r\n"),
])
def test_parser_hashes_original_bytes_separately_from_normalized_text(filename, raw):
    from deep_research.workbench.datasets import parse_table_file

    parsed = parse_table_file(raw, filename)
    assert parsed["file_sha256"] == hashlib.sha256(raw).hexdigest()
    sheet = parsed["sheets"][0]
    assert sheet["input_sha256"] == _digest(sheet["csv"])
    assert sheet["input_sha256"] != parsed["file_sha256"]


async def test_uploaded_xlsx_hashes_survive_creation_freeze_and_export(api_repo):
    api, repo = api_repo
    raw = _xlsx({"Data": [["x", "y"], [1, 2], [2, None], [3, 4]], "Other": [["z"], [8]]})
    async with _client(api.app) as client:
        uploaded = await client.post("/api/datasets", json={
            "filename": "study.xlsx", "data_base64": base64.b64encode(raw).decode(),
        })
        parsed = uploaded.json()
        assert parsed["file_sha256"] == hashlib.sha256(raw).hexdigest()
        sheet = parsed["sheets"][0]
        created = await client.post("/api/runs", json={
            "query": "描述测量值", "template": "dataAnalysis", "clarified": True,
            "dataset": sheet["csv"], "dataset_source": {
                "filename": parsed["filename"], "sheet": sheet["name"],
                "file_sha256": parsed["file_sha256"], "input_sha256": sheet["input_sha256"],
            },
        })
    assert created.status_code == 202, created.text
    contract = await _contract(repo, created.json()["run_id"])
    result = analyse(contract["dataset_csv"], source=contract["dataset_source"])
    frozen = json.loads(json.dumps(result.snapshot()))
    restored = analyse(contract["dataset_csv"], frozen=frozen)
    workbook = load_workbook(io.BytesIO(_stats_xlsx(restored)))
    overview = _overview(workbook)
    assert overview["原始文件字节 SHA-256"] == hashlib.sha256(raw).hexdigest()
    assert overview["分析输入文本 SHA-256"] == _digest(sheet["csv"])
    assert overview["原始文件字节 SHA-256"] != overview["分析输入文本 SHA-256"]
    assert restored.snapshot() == frozen


@pytest.mark.parametrize("merged", [False, True])
async def test_creation_rejects_stale_input_hash_before_creating_a_run(api_repo, merged):
    from tests.test_dataset_merge import merge_request

    api, repo = api_repo
    body = {"query": "描述测量值", "template": "dataAnalysis", "clarified": True,
            "dataset_source": {"filename": "study.xlsx", "input_sha256": "0" * 64}}
    body.update({"dataset_merge": merge_request()} if merged else {"dataset": "x\n1\n2"})
    async with _client(api.app) as client:
        response = await client.post("/api/runs", json=body)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "dataset_source_mismatch"
    assert not await repo.list_runs(limit=10)


def test_old_source_requests_keep_their_original_idempotency_hash():
    from deep_research.api import CreateRunRequest, _run_request_hash

    request = CreateRunRequest(
        query="test", dataset="x\n1\n2", dataset_source={"filename": "a.csv"},
    )
    old = request.model_dump(mode="json")
    old.pop("dataset_merge")
    old["dataset_source"] = {"filename": "a.csv", "sheet": ""}
    expected = hashlib.sha256(
        json.dumps(old, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    assert _run_request_hash(request) == expected
    changed = request.model_copy(update={"dataset_source": request.dataset_source.model_copy(
        update={"file_sha256": "a" * 64}
    )})
    assert _run_request_hash(changed) != expected


@pytest.mark.parametrize("field", ["file_sha256", "input_sha256"])
def test_source_hash_requires_a_complete_sha256_digest(field):
    from pydantic import ValidationError

    from deep_research.api import DatasetSource

    with pytest.raises(ValidationError):
        DatasetSource.model_validate({field: "invalid"})


def test_merge_workbook_includes_source_table_hashes_and_join_counts():
    from deep_research.workbench.dataset_merge import DatasetMerge, merge_tables
    from tests.test_dataset_merge import merge_request

    request = merge_request()
    merged = merge_tables(DatasetMerge.model_validate(request))
    result = analyse(merged["csv"], source={
        "filename": "study.xlsx", "file_sha256": "a" * 64, "merge": merged["merge"],
    })
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    assert list(workbook["合并来源"].values)[1:] == [
        (table["name"], 3, _digest(table["csv"])) for table in request["tables"]
    ]
    rows = list(workbook["连接计划"].values)
    joined = dict(zip(rows[0], rows[1], strict=True))
    assert joined["主表键"] == '["subject_id"]' and joined["右表键"] == '["code"]'
    assert joined["主表未匹配行数"] == joined["右表未匹配行数"] == 1
    assert joined["连接方式"] == "left" and joined["关系"] == "one_to_one"
    assert _overview(workbook)["合并主表"] == "Measurements"
    assert _overview(workbook)["分析输入文本 SHA-256"] == _digest(merged["csv"])
