"""Regression cases for verified design defects across data, revision and delivery."""

from __future__ import annotations

import asyncio
import io
from copy import deepcopy
from types import SimpleNamespace

import pytest
from pptx import Presentation

from deep_research.agents.base import Blackboard, RunContext
from deep_research.models import Report, ResearchResult
from deep_research.observability import Tracer
from deep_research.orchestrator import create_initial_execution
from deep_research.persistence.repository import RunDetail
from deep_research.workbench.analysis import DatasetError, analyse
from deep_research.workbench.contract import CONTRACT_SCRATCH_KEY, build_contract
from deep_research.workbench.delivery.pptx import fit_report, paginate_deck, render_pptx
from deep_research.workbench.publish import (
    DeliveryBundle,
    DeliveryFile,
    _stats_xlsx,
    build_bundle,
    delivery_fingerprint,
)
from deep_research.workbench.quality import QualityPolicy
from deep_research.workbench.revision import Assessment, write_with_revisions
from deep_research.workbench.scholarly import check_sources
from deep_research.workbench.templates import get_template
from deep_research.workbench.writers import TemplateWriter
from tests.fakes import FakeLLM, FakeSearch, verified_finding


def test_pairwise_missingness_direction_and_confidence_interval():
    result = analyse("scene,a,b\n1,10,12\n2,20,21\n3,30,34\n4,40,\n5,,99", "请做配对差异分析")
    test = next(item for item in result.tests if item.get("paired"))
    assert test["n_pairs"] == 3 and test["excluded_pairs"] == 2
    assert test["left"] == "a" and test["right"] == "b"
    assert test["mean_difference"] == pytest.approx(7 / 3, abs=0.0001)
    assert test["statistic"] == pytest.approx(2.6458, abs=0.0001)
    assert test["p_value"] == pytest.approx(0.1181, abs=0.0001)
    assert test["df"] == 2 and test["ci_low"] < 0 < test["ci_high"]
    assert test["significant"] is False
    assert "scene" not in result.numeric and "95%" in result.facts()
    assert any("paired" in figure.name for figure in result.figures)
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    headers = [cell.value for cell in workbook["显著性检验"][1]]
    assert {"n_pairs", "excluded_pairs", "ci_low", "ci_high"} <= set(headers)


@pytest.mark.parametrize("rows", ["1,10,11\n2,20,21", "1,10,11\n2,,21"])
def test_degenerate_pairs_are_unavailable_not_nonsignificant(rows):
    result = analyse("scene,a,b\n" + rows, "配对比较")
    test = next(item for item in result.tests if item.get("paired"))
    assert test["significant"] is None and test["p_value"] == "NA"
    assert result.issues and "无法检验" in result.facts()


def test_integer_measurement_columns_are_not_misclassified_as_identifiers():
    result = analyse("id,psnr\n1,30\n2,31\n3,30\n4,31", "描述统计")
    assert result.numeric == ["psnr"]
    assert result.describe[0]["n"] == 4 and result.describe[0]["mean"] == 30.5
    assert not result.tests


def test_nonfinite_input_is_rejected_and_ambiguous_pairing_is_not_guessed():
    with pytest.raises(DatasetError, match="无穷值"):
        analyse("x,y\n1,inf\n2,3")
    result = analyse("a,b,c\n1,3,5\n2,4,6\n3,4,7", "做配对检验")
    assert not any(test.get("paired") for test in result.tests)
    assert any("不明确" in issue for issue in result.issues)


def test_paired_long_table_does_not_fall_back_to_an_independent_test():
    result = analyse("subject,method,psnr\n1,A,30\n1,B,31\n2,A,32\n2,B,34", "做配对检验")
    assert result.numeric == ["psnr"]
    assert result.tests == [] and result.issues
    independent = analyse(
        "subject,method,psnr\n1,A,30\n1,B,31\n2,A,32\n2,B,34", "非配对独立样本检验"
    )
    assert not any(test.get("paired") for test in independent.tests)


def test_download_uses_frozen_statistics_and_rejects_changed_data():
    csv = "scene,a,b\n1,10,12\n2,20,21\n3,30,34"
    old = analyse(csv, "只做描述统计")
    restored = analyse(csv, "请做配对检验", frozen=old.snapshot())
    assert restored.tests == old.tests == []
    assert restored.describe == old.describe
    assert [f.name for f in restored.figures] == [f.name for f in old.figures]
    with pytest.raises(DatasetError, match="快照不一致"):
        analyse(csv.replace("34", "35"), frozen=old.snapshot())


async def test_long_revision_preserves_the_tail_in_the_actual_retry_request():
    previous = "前半部分正文。" * 2500 + "\n\n结尾关键结论：TAIL-KEEP"
    calls = 0

    async def write(revision):
        nonlocal calls
        calls += 1
        if revision is None:
            return previous
        assert "TAIL-KEEP" in revision
        return previous + "\n修订完成"

    body, _ = await write_with_revisions(
        write,
        lambda body: Assessment(hard=[] if body.endswith("修订完成") else ["需要修订"]),
        max_revisions=1,
    )
    assert calls == 2 and "TAIL-KEEP" in body


async def test_best_draft_restores_its_structured_artifact_instead_of_last_attempt(
    settings, monkeypatch
):
    from deep_research.workbench import writers

    class Versions(TemplateWriter):
        output_keys = ("_slide_deck",)
        count = 0

        async def write(self, bb, ctx, template, contract, material, revision=None):
            self.count += 1
            bb.scratch["_slide_deck"] = {"version": self.count}
            return f"body-{self.count}"

    monkeypatch.setattr(
        writers, "assess_draft", lambda body, **_: Assessment(hard=["x"] * int(body[-1]))
    )
    bb = Blackboard(query="q")
    ctx = RunContext(llm=FakeLLM(), search_tool=FakeSearch(), tracer=Tracer(), settings=settings)
    body, log = await Versions()._write_checked(
        bb,
        ctx,
        get_template("slides"),
        None,
        "",
        {},
        policy=QualityPolicy(max_revisions=1),
        min_citations=0,
        require_corroboration=False,
    )
    assert body == "body-1" and log.chosen == 1
    assert bb.scratch["_slide_deck"] == {"version": 1}


async def test_mindmap_own_revision_path_also_restores_the_selected_version(settings):
    from deep_research.workbench.writers import Mindmap, MindmapWriter, mindmap_to_markdown

    class Versions(MindmapWriter):
        count = 0

        async def write(self, bb, *args):
            self.count += 1
            raw = {
                "root": "主题",
                "branches": [
                    {
                        "label": "分支",
                        "children": [{"label": "节点"} for _ in range(4 if self.count == 1 else 1)],
                    }
                    for _ in range(6 if self.count == 1 else 1)
                ],
            }
            bb.scratch["_mindmap"] = raw
            return mindmap_to_markdown(Mindmap.model_validate(raw))

    settings.quality = {"max_revisions": 1}
    ctx = RunContext(llm=FakeLLM(), search_tool=FakeSearch(), tracer=Tracer(), settings=settings)
    result = await Versions().step(Blackboard(query="概念结构"), ctx)
    extras = result.scratch["workbench"]["extras"]
    assert extras["revision"]["chosen"] == 1
    assert extras["stats"]["branches"] == 6


def test_pptx_continuation_preserves_every_bullet_note_and_citation():
    bullets = [
        "长论述。" * 60,
        "multi\nline\ntext\nwith\nmore\nlines",
        *[f"要点 {i}" for i in range(8)],
    ]
    deck = {
        "title": "汇报",
        "slides": [{"title": "内容", "bullets": bullets, "notes": "完整备注", "citations": [1]}],
    }
    before = deepcopy(deck)
    pages = paginate_deck(deck)
    assert deck == before
    assert "".join(b for page in pages["slides"] for b in page["bullets"]) == "".join(bullets)
    assert not fit_report(pages)
    assert all(p["notes"] == "完整备注" and p["citations"] == [1] for p in pages["slides"])
    presentation = Presentation(io.BytesIO(render_pptx(deck)))
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "要点 7" in text and len(presentation.slides) > 2


def _detail(settings):
    template = get_template("slides")
    execution = create_initial_execution("q", template.workflow, settings)
    execution.checkpoint["scratch"][CONTRACT_SCRATCH_KEY] = build_contract(
        template, "q"
    ).model_dump(mode="json")
    execution.checkpoint["scratch"]["workbench"] = {
        "template": "slides",
        "extras": {
            "deck": {
                "title": "旧幻灯片",
                "slides": [
                    {
                        "title": "旧结果",
                        "bullets": ["伪造值 999.9"],
                        "notes": "旧备注",
                        "citations": [1],
                    }
                ],
            }
        },
    }
    return RunDetail(
        id="run",
        query="q",
        status="done",
        report=Report(
            query="q", markdown="## 已验证素材摘要\n\n- 发现X [1]\n", citations=["https://a.com"]
        ),
        results=[ResearchResult(sub_question="q", findings=[verified_finding()])],
        orchestration=execution,
    )


def test_pptx_does_not_reintroduce_an_unreviewed_deck_after_report_fallback(settings):
    bundle = build_bundle(_detail(settings))
    pptx = next(file for file in bundle.files if file.format == "pptx")
    presentation = Presentation(io.BytesIO(pptx.data))
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "999.9" not in text and "发现X" in text
    assert pptx.status == "warn" and any(g.name == "structured_content" for g in bundle.gates)


def test_delivery_cache_fingerprint_covers_non_markdown_inputs(settings):
    detail = _detail(settings)
    initial = delivery_fingerprint(detail)
    detail.orchestration.checkpoint["scratch"]["workbench"]["extras"]["deck"]["title"] = "修订标题"
    assert delivery_fingerprint(detail) != initial
    changed = delivery_fingerprint(detail)
    detail.orchestration.checkpoint["scratch"][CONTRACT_SCRATCH_KEY]["dataset_csv"] = "a\n1"
    assert delivery_fingerprint(detail) != changed


async def test_delivery_generation_survives_first_waiter_disconnect(settings, monkeypatch):
    from deep_research.workbench import api as workbench_api

    detail = _detail(settings)
    entered, release = asyncio.Event(), asyncio.Event()
    calls = 0
    bundle = DeliveryBundle("slides", "report", [], [], "pass", "")

    async def get_run(_):
        return detail

    async def render(*args):
        nonlocal calls
        calls += 1
        entered.set()
        await release.wait()
        return bundle

    monkeypatch.setattr(workbench_api, "run_blocking", render)
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(repo=SimpleNamespace(get_run=get_run)))
    )
    first = asyncio.create_task(workbench_api._bundle(request, "run"))
    await entered.wait()
    first.cancel()
    with pytest.raises(asyncio.CancelledError):
        await first
    second = asyncio.create_task(workbench_api._bundle(request, "run"))
    release.set()
    assert await second is bundle
    assert await workbench_api._bundle(request, "run") is bundle
    assert calls == 1


def test_failed_files_are_not_selected_as_primary_delivery():
    failed = DeliveryFile("bad.pdf", "pdf", "bad", "report", b"bad", status="fail")
    usable = DeliveryFile("ok.html", "html", "ok", "reading", b"ok")
    bundle = DeliveryBundle("t", "t", [failed, usable], [], "fail", "")
    assert bundle.registry()["primary"] == "ok.html"
    bundle.files = [failed]
    assert bundle.registry()["primary"] is None


def test_document_identity_retains_meaningful_query_parameters():
    urls = ["https://publisher.org/article?id=A", "https://publisher.org/article?id=a"]
    assert not check_sources(urls, 2, 2)
    chunks = [
        "https://workspace.invalid/attachments/paper?chunk=1",
        "https://workspace.invalid/attachments/paper?chunk=2",
    ]
    assert any(item.code == "duplicate-source" for item in check_sources(chunks, 2, 2))


def test_numeric_verification_distinguishes_opposite_scientific_exponents():
    from deep_research.report.validation import validate_body

    finding = verified_finding(statement="p=1e-5", evidence_quote="p=1e-5")
    results = [ResearchResult(sub_question="q", findings=[finding])]
    mapping = {finding.source_url: 1}
    assert not validate_body("p=1e−5 [1]。", results, mapping).issues
    assert "unsupported_number" in validate_body("p=1e+5 [1]。", results, mapping).issues
