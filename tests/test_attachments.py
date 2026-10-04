"""任务附件回归：上传解析（PDF / Word / PPT / Excel / 文本）、接口、模型阅读并逐字核验。"""

from __future__ import annotations

import base64
import io
import tempfile

import httpx
import pytest
from httpx import ASGITransport

from deep_research.models import ExtractedFindingList
from deep_research.workbench.attachments import (
    ATTACHMENT_URL_PREFIX,
    ATTACHMENTS_SCRATCH_KEY,
    AttachmentError,
    limit_attachments,
    parse_attachment,
)
from tests.fakes import FakeLLM, FakeSearch


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("实验结果", 1)
    document.add_paragraph("在 CAVE 数据集上，重建 PSNR 达到 38.4 dB，优于基线 1.2 dB。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "方法", "PSNR"
    table.cell(1, 0).text, table.cell(1, 1).text = "Ours", "38.4"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "方法概述"
    slide.placeholders[1].text = "深度展开网络将迭代优化映射为可学习层"
    slide.notes_slide.notes_text_frame.text = "强调物理先验"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "结果"
    sheet.append(["method", "psnr"])
    sheet.append(["A", 30.1])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _pdf() -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Snapshot spectral imaging reconstruction reaches 38.4 dB PSNR.")
    data = document.tobytes()
    document.close()
    return bytes(data)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "builder", "kind", "expected", "locator"),
    [
        ("实验记录.docx", _docx, "docx", "38.4 dB", "实验结果"),
        ("组会.pptx", _pptx, "pptx", "可学习层", "第 1 张幻灯片"),
        ("结果.xlsx", _xlsx, "xlsx", "30.1", "工作表 结果"),
        ("paper.pdf", _pdf, "pdf", "38.4 dB", ""),
    ],
)
async def test_parse_supported_formats_with_locators(filename, builder, kind, expected, locator):  # type: ignore[no-untyped-def]
    raw = builder()
    attachment = await parse_attachment(raw, filename)
    assert attachment.kind == kind
    text = "\n".join(chunk.content for chunk in attachment.chunks)
    assert expected in text
    assert locator in attachment.chunks[0].locator
    source = attachment.sources()[0]
    assert source.url.startswith(ATTACHMENT_URL_PREFIX) and source.title == filename
    # 同一文件重复上传得到同一 id（内容摘要）
    assert (await parse_attachment(raw, filename)).id == attachment.id


@pytest.mark.asyncio
async def test_parse_text_and_reject_unsupported_or_empty():  # type: ignore[no-untyped-def]
    note = await parse_attachment("# 笔记\n\n快照光谱成像的误差来源。".encode(), "note.md")
    assert note.kind == "markdown" and "误差来源" in note.chunks[0].content
    with pytest.raises(AttachmentError, match="不支持"):
        await parse_attachment(b"MZ\x00\x00", "tool.exe")
    with pytest.raises(AttachmentError, match="为空"):
        await parse_attachment(b"", "empty.txt")


async def test_pdf_metadata_reaches_bibliography_only_when_confirmed_on_first_page():
    import hashlib

    import pymupdf

    from deep_research.bibliography import build_bibliography
    from deep_research.guardrails import EvidenceVerifier
    from deep_research.models import Finding

    title = "Deep Spectral Reconstruction"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), title + "\nAlice Smith and Bob Jones\nVerified method description.")
    document.set_metadata(
        {
            "title": title,
            "author": "Alice Smith; Bob Jones; Unknown Person",
            "creationDate": "D:20260101000000",
            "subject": "Imaginary Journal",
        }
    )
    raw = document.tobytes()
    document.close()
    attachment = await parse_attachment(raw, "short.pdf")
    assert attachment.id == hashlib.sha256(raw).hexdigest()[:24]
    assert attachment.filename == "short.pdf" and attachment.title == title
    assert attachment.authors == ["Alice Smith", "Bob Jones"]
    source = attachment.sources()[0]
    assert source.scholarly is None  # A PDF author is not evidence of publication status.
    finding = Finding(
        statement="Verified method description.",
        source_url=source.url,
        evidence_quote="Verified method description.",
    )
    finding = EvidenceVerifier().verify(finding, source).finding
    assert finding is not None
    catalog = build_bibliography("Method [1].", [source.url], [finding], [source])
    assert (
        catalog.documents[0].reference
        == "Alice Smith, Bob Jones. " + title + ". 年份未识别. 出处未识别"
    )
    assert catalog.documents[0].url == ""


@pytest.mark.parametrize("author", ["Alice Smith; Bob Jones", "Alice Smith, Bob Jones"])
def test_pdf_author_delimiters_and_stale_template_metadata(author):
    from deep_research.tools.oa_pdf_fulltext import _confirmed_metadata

    first_page = "Deep Spectral Reconstruction\nAlice Smith and Bob Jones"
    assert _confirmed_metadata({"title": "Old unrelated report", "author": author}, first_page) == (
        "",
        (),
    )
    assert _confirmed_metadata(
        {"title": "Deep Spectral Reconstruction", "author": author}, first_page
    ) == (
        "Deep Spectral Reconstruction",
        ("Alice Smith", "Bob Jones"),
    )


@pytest.mark.asyncio
async def test_limit_attachments_dedupes_and_bounds_chunks():  # type: ignore[no-untyped-def]
    note = await parse_attachment("正文".encode(), "a.txt")
    kept = limit_attachments([note, note])
    assert len(kept) == 1


async def test_long_attachments_keep_all_chunks_and_late_files():
    files = [
        await parse_attachment(
            (f"Paper {i}\n" + "Continuous text. " * 12000 + f"END-{i}").encode(), f"{i}.txt"
        )
        for i in range(3)
    ]
    assert all(len(item.chunks) > 40 and not item.truncated for item in files)
    assert sum(len(item.chunks) for item in files) > 120
    kept = limit_attachments(files)
    assert [item.model_dump() for item in kept] == [item.model_dump() for item in files]
    assert all(f"END-{i}" in item.chunks[-1].content for i, item in enumerate(kept))


async def test_attachment_bounds_reject_instead_of_changing_material(monkeypatch):
    from deep_research.workbench import attachments

    files = [await parse_attachment(f"Paper {i} body".encode(), f"{i}.txt") for i in range(9)]
    with pytest.raises(AttachmentError, match="最多提交"):
        limit_attachments(files)
    with pytest.raises(AttachmentError, match="未完整"):
        limit_attachments([files[0].model_copy(update={"truncated": True})])
    monkeypatch.setattr(attachments, "MAX_TOTAL_PARSED_CHARS", 15)
    with pytest.raises(AttachmentError, match="合计文本"):
        limit_attachments(files[:2])
    assert all(not item.truncated and len(item.chunks) == 1 for item in files)


def _client(app) -> httpx.AsyncClient:  # type: ignore[no-untyped-def]
    return httpx.AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


@pytest.mark.asyncio
async def test_upload_endpoint_and_run_creation_freeze_attachments(monkeypatch):  # type: ignore[no-untyped-def]
    import asyncio

    from deep_research import api
    from deep_research.config import Settings
    from deep_research.persistence.memory_repository import InMemoryRepository

    repo = InMemoryRepository()
    monkeypatch.setattr(api.app.state, "catalog", None, raising=False)
    monkeypatch.setattr(api, "_run_limiter", api._RateLimiter(1000, 60), raising=False)
    api.app.state.settings = Settings()
    api.app.state.repo = repo
    api.app.state.live, api.app.state.tasks = {}, set()
    api.app.state.run_tasks, api.app.state.cancellation_requested = {}, set()
    api.app.state.config_lock = asyncio.Lock()
    payload = {"filename": "实验记录.docx", "data_base64": base64.b64encode(_docx()).decode()}
    async with _client(api.app) as client:
        uploaded = await client.post("/api/attachments", json=payload)
        bad = await client.post(
            "/api/attachments", json={"filename": "x.exe", "data_base64": "TVo="}
        )
        body = uploaded.json()
        created = await client.post(
            "/api/runs",
            json={
                "query": "总结实验结果",
                "template": "autoResearch",
                "strategy": "quick",
                "clarified": True,
                "attachments": [body["attachment"]],
            },
        )
    assert uploaded.status_code == 201, uploaded.text
    assert body["summary"]["kind"] == "docx" and body["summary"]["chunk_count"] >= 1
    assert bad.status_code == 422
    assert created.status_code == 202, created.text
    detail = await repo.get_run(created.json()["run_id"])
    assert detail is not None and detail.orchestration is not None
    frozen = detail.orchestration.checkpoint["scratch"][ATTACHMENTS_SCRATCH_KEY]
    assert frozen[0]["filename"] == "实验记录.docx"


class AttachmentLLM(FakeLLM):
    """抽取时引用附件片段 URL 与其中逐字出现的原文。"""

    async def parse(self, system, user, schema, *, temperature=0.2, retries=2):  # type: ignore[no-untyped-def]
        if schema is ExtractedFindingList and ATTACHMENT_URL_PREFIX in user:
            start = user.index(ATTACHMENT_URL_PREFIX)
            url = user[start:].split()[0].rstrip("）)]，,")
            return ExtractedFindingList(
                findings=[
                    {
                        "statement": "方法在 CAVE 数据集上达到 38.4 dB PSNR",
                        "source_url": url,
                        "evidence_quote": "重建 PSNR 达到 38.4 dB",
                        "confidence": 0.9,
                    }
                ]
            )
        return await super().parse(system, user, schema, temperature=temperature, retries=retries)


@pytest.mark.asyncio
async def test_model_reads_attachment_and_cites_it_after_verification():  # type: ignore[no-untyped-def]
    from deep_research.config import Settings
    from deep_research.orchestrator import DeepResearchAgent, create_initial_execution
    from deep_research.persistence.memory_repository import InMemoryRepository
    from deep_research.workbench.contract import CONTRACT_SCRATCH_KEY, build_contract
    from deep_research.workbench.templates import get_template

    attachment = await parse_attachment(_docx(), "实验记录.docx")
    settings = Settings(artifact_root=tempfile.mkdtemp())
    settings.max_rounds = 0
    template = get_template("autoResearch")
    assert template is not None
    workflow = template.workflow_for("quick")
    execution = create_initial_execution("总结实验结果", workflow, settings)
    scratch = execution.checkpoint.setdefault("scratch", {})
    scratch[CONTRACT_SCRATCH_KEY] = build_contract(template, "总结实验结果").model_dump(mode="json")
    scratch[ATTACHMENTS_SCRATCH_KEY] = [attachment.model_dump(mode="json")]
    repo = InMemoryRepository()
    run_id = await repo.create_run("总结实验结果", execution=execution)
    agent = DeepResearchAgent(
        settings,
        llm=AttachmentLLM(),
        search_tool=FakeSearch(),
        workflow=workflow,
        repo=repo,
        run_id=run_id,
        initial_execution=execution,
    )
    await agent.run("总结实验结果")
    detail = await repo.get_run(run_id)
    assert detail is not None
    verified = [
        f
        for r in detail.results
        for f in r.findings
        if f.source_url.startswith(ATTACHMENT_URL_PREFIX) and f.verification.status == "verified"
    ]
    assert verified, "附件中的原文必须能通过逐字核验"
    assert verified[0].verification.source_title == "实验记录.docx"
    # 参考来源显示文件名与定位，而不是内部占位 URL
    assert verified[0].verification.source_reference.startswith("实验记录.docx（实验结果")


class NoSearch(FakeSearch):
    async def search(self, query, *, max_results=5):  # type: ignore[no-untyped-def]
        raise AssertionError("paper tasks with uploaded papers must not open-search")


@pytest.mark.asyncio
async def test_paper_read_with_only_an_uploaded_paper_reads_it_without_searching():  # type: ignore[no-untyped-def]
    """只上传论文、没给链接时：以附件为精读对象，不退化成按主题检索别的论文。"""
    from deep_research.config import Settings
    from deep_research.orchestrator import DeepResearchAgent, create_initial_execution
    from deep_research.persistence.memory_repository import InMemoryRepository
    from deep_research.workbench.contract import CONTRACT_SCRATCH_KEY, build_contract
    from deep_research.workbench.delivery_store import load_version
    from deep_research.workbench.intake import INTAKE_SOURCES_KEY
    from deep_research.workbench.templates import get_template

    attachment = await parse_attachment(_docx(), "待精读论文.docx")
    settings = Settings(artifact_root=tempfile.mkdtemp())
    template = get_template("paperRead")
    assert template is not None
    query = "重点看实验部分"
    execution = create_initial_execution(query, template.workflow, settings)
    scratch = execution.checkpoint.setdefault("scratch", {})
    scratch[CONTRACT_SCRATCH_KEY] = build_contract(template, query).model_dump(mode="json")
    scratch[ATTACHMENTS_SCRATCH_KEY] = [attachment.model_dump(mode="json")]
    repo = InMemoryRepository()
    run_id = await repo.create_run(query, execution=execution)
    agent = DeepResearchAgent(
        settings,
        llm=AttachmentLLM(),
        search_tool=NoSearch(),
        workflow=template.workflow,
        repo=repo,
        run_id=run_id,
        initial_execution=execution,
    )
    await agent.run(query)
    detail = await repo.get_run(run_id)
    assert detail is not None
    completion = detail.orchestration.checkpoint["scratch"].get("_completion")
    assert isinstance(completion, dict), detail.status
    record = load_version(detail, settings.artifact_root, completion["content_version"]).registry()
    required = set(template.deliverables)
    missing = required - {item["format"] for item in record["items"]}
    nonpassing = [gate for gate in record["gates"] if gate["status"] != "pass"]
    requires_review = (
        bool(missing or nonpassing or record["failures"])
        or record["status"] != "pass"
        or any(
            item["status"] != "pass" or item["size"] <= 0
            for item in record["items"]
            if item["format"] in required
        )
    )
    expected = "needs_review" if requires_review else "done"
    # Reading and verifying the supplied original succeeds; the fixture's short
    # generic draft still lacks the sections promised by a full paper-reading task.
    assert {"length", "structure"} <= {gate["name"] for gate in nonpassing}
    assert expected == "needs_review"
    assert detail.status == completion["status"] == expected
    assert completion["required_formats"] == sorted(required)
    assert completion["content_version"] == record["content_version"]
    assert completion["input_version"] == record["input_version"]
    assert completion["gates"] == record["gates"]
    assert bool(completion["issues"]) is requires_review
    for gate in nonpassing:
        assert set(gate["issues"]) <= set(completion["issues"])
    for fmt in missing:
        assert any(fmt.upper() in issue for issue in completion["issues"])
    verified = [
        f
        for r in detail.results
        for f in r.findings
        if f.source_url.startswith(ATTACHMENT_URL_PREFIX) and f.verification.status == "verified"
    ]
    assert verified, "精读对象就是上传的论文，其原文须能通过逐字核验"
    urls = [f.source_url for r in detail.results for f in r.findings]
    assert all(url.startswith(ATTACHMENT_URL_PREFIX) for url in urls)
    intake = detail.orchestration.checkpoint["scratch"][INTAKE_SOURCES_KEY]
    assert intake["mode"] == "attachments"
    assert intake["sections"][0]["title"] == "待精读论文.docx"
