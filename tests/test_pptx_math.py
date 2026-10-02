from __future__ import annotations

import io
from copy import deepcopy
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from pptx import Presentation

from deep_research.workbench.delivery.math import MathRenderError
from deep_research.workbench.delivery.pptx import (
    _estimated_lines,
    _math_spans,
    paginate_deck,
    render_pptx,
)

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"


def xml_slides(data):
    with ZipFile(io.BytesIO(data)) as z:
        assert z.testzip() is None
        return [
            ET.fromstring(z.read(name))
            for name in z.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]


def test_pptx_uses_editable_math_and_rich_runs_with_no_visible_tex_commands():
    deck = {
        "title": "公式汇报",
        "slides": [
            {
                "title": r"方法 $x_i$",
                "bullets": [
                    r"**权重**为 $\sigma_j$，矩阵乘积为 $K_j^T Q_j$，误差为 $\frac{a}{b}$。",
                    r"价格 \$10K–\$100K，代码 `x_i` 保留。",
                ],
                "notes": "完整备注",
                "citations": [1],
            }
        ],
    }
    before = deepcopy(deck)
    data = render_pptx(deck)
    assert deck == before
    roots = xml_slides(data)
    maths = [node for root in roots for node in root.findall(f".//{A14}m")]
    assert len(maths) == 4
    assert any(root.findall(f".//{M}f") for root in roots)
    assert any(root.findall(f".//{M}sSub") for root in roots)
    text = "".join(node.text or "" for root in roots for node in root.findall(f".//{A}t"))
    assert "\\sigma" not in text and "$K" not in text and "**权重**" not in text
    assert "$10K–$100K" in text and "x_i" in text
    assert any(node.attrib.get("b") == "1" for root in roots for node in root.findall(f".//{A}rPr"))
    assert Presentation(io.BytesIO(data)).slides[1].notes_slide.notes_text_frame.text == "完整备注"


@pytest.mark.parametrize("first_title", ["同一标题", "标题页", "封面", "Title slide"])
def test_existing_title_slide_does_not_gain_a_second_cover_or_lose_content(first_title):
    deck = {
        "title": "同一标题",
        "subtitle": "十二分钟汇报",
        "slides": [
            {
                "title": first_title,
                "bullets": ["作者与研究对象"],
                "notes": "开场备注",
                "citations": [1],
            },
            {"title": "结论", "bullets": ["核心结论"], "notes": "结论备注", "citations": [1]},
        ],
    }
    data = render_pptx(deck, citations=["原始文献"])
    pres = Presentation(io.BytesIO(data))
    assert len(pres.slides) == 3
    text = "\n".join(s.text for s in pres.slides[0].shapes if s.has_text_frame)
    assert "同一标题" in text and "十二分钟汇报" in text and "作者与研究对象" in text
    assert pres.slides[0].notes_slide.notes_text_frame.text == "开场备注"


def test_pagination_preserves_formula_boundaries_and_all_raw_content():
    bullet = (
        "普通描述。" * 70 + r" 连续公式 $\frac{a + b}{c + d}$ 及 $x_i^2$ 保留。" + "后续描述。" * 40
    )
    deck = {"title": "汇报", "slides": [{"title": "长内容", "bullets": [bullet], "notes": "备注"}]}
    pages = paginate_deck(deck)
    parts = [b for page in pages["slides"] for b in page["bullets"]]
    assert "".join(parts) == bullet
    assert sum(len(_math_spans(part)) for part in parts) == 2
    roots = xml_slides(render_pptx(deck))
    assert sum(len(root.findall(f".//{M}oMath")) for root in roots) == 2


def test_one_item_continuation_is_balanced_without_reordering_or_losing_text():
    bullets = [f"内容{i}：" + "测" * count for i, count in enumerate([40, 40, 72, 72, 40])]
    deck = {
        "title": "汇报",
        "slides": [{"title": "内容", "bullets": bullets, "notes": "备注", "citations": [1]}],
    }
    pages = paginate_deck(deck)["slides"]
    assert len(pages) == 2 and all(len(page["bullets"]) >= 2 for page in pages)
    assert [text for page in pages for text in page["bullets"]] == bullets
    assert all(page["notes"] == "备注" and page["citations"] == [1] for page in pages)


def test_math_rules_do_not_convert_literal_prices_or_code_and_handle_brackets():
    text = r"价格 \$5 和 \$10，代码 `$x_i$`；公式 \(x_i\) 与 $y_j$。"
    spans = _math_spans(text)
    assert [tex for _, tex, _ in spans.values()] == ["x_i", "y_j"]
    assert _estimated_lines(r"$\frac{1}{2}$") >= 1


def test_orphan_script_has_no_office_placeholder_and_east_asian_font_is_explicit():
    roots = xml_slides(
        render_pptx(
            {
                "title": "中文标题",
                "slides": [{"title": "方法", "bullets": [r"head$_j$ 内的注意力"]}],
            }
        )
    )
    bases = [e for root in roots for e in root.findall(f".//{M}sSub/{M}e")]
    assert bases and all("\u200b" in "".join(base.itertext()) for base in bases)
    assert all(
        node.attrib.get("typeface") == "Microsoft YaHei"
        for root in roots
        for node in root.findall(f".//{A}ea")
    )


@pytest.mark.parametrize("tex", [r"\input{secret}", r"\madeup{x}"])
def test_invalid_math_fails_instead_of_being_exported_as_raw_code(tex):
    with pytest.raises(MathRenderError):
        render_pptx({"title": "公式", "slides": [{"title": "方法", "bullets": [f"${tex}$"]}]})
