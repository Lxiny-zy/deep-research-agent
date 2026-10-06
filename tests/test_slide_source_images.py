"""Task-owned PDF page registration and the actual immutable image/PPTX boundary."""

import hashlib
import io
from copy import deepcopy
from dataclasses import replace

import fitz
import pytest
from PIL import Image
from pptx import Presentation

from deep_research.models import ResearchResult
from deep_research.workbench.attachments import attachment_url, parse_attachment, save_original
from deep_research.workbench.delivery_render import render_bundle
from deep_research.workbench.publish import DeliveryFile
from deep_research.workbench.slide_content import SlideDeck, compile_deck
from deep_research.workbench.slide_images import (
    freeze_selected_images,
    render_registered_page,
    source_image_catalog,
)
from deep_research.workbench.support import evidence_id
from tests.fakes import verified_finding


async def source_fixture(tmp_path, settings):
    with fitz.open() as pdf:
        page = pdf.new_page(width=420, height=500)
        page.insert_text((40, 50), "Figure 1. Experimental apparatus.")
        page.draw_rect(fitz.Rect(40, 90, 360, 380), color=(1, 0, 0), fill=(0.8, 0.9, 1))
        page.insert_text((40, 420), "Original caption and surrounding context.")
        raw = pdf.tobytes()
    settings = replace(settings, artifact_root=str(tmp_path))
    attachment = await parse_attachment(raw, "experiment.pdf")
    save_original(settings, attachment.id, raw)
    attachment.stored = True
    source_url = attachment_url(attachment.id, attachment.chunks[0].ordinal)
    finding = verified_finding(
        "Figure 1 shows the experimental apparatus.",
        source_url,
        "Figure 1. Experimental apparatus.",
    )
    results = [ResearchResult(sub_question="Apparatus", findings=[finding])]
    scratch = {"attachments": [attachment.model_dump(mode="json")]}
    deck = SlideDeck.model_validate(
        {
            "title": "Experiment",
            "slides": [
                {
                    "title": "Apparatus",
                    "bullets": [],
                    "notes": "Explain the original apparatus and its context.",
                    "citations": [1],
                    "image": {"finding_id": evidence_id(finding), "figure_label": "Figure 1"},
                }
            ],
        }
    )
    return settings, deck, results, scratch, source_url


@pytest.mark.parametrize("region", [False, True])
async def test_original_page_is_frozen_and_delivered_as_actual_image(tmp_path, settings, region):
    settings, deck, results, scratch, url = await source_fixture(tmp_path, settings)
    if region:
        scratch["attachments"][0]["image_regions"] = [
            {"page": 1, "figure_label": "Fig. 1", "bounds": [0.1, 0.18, 0.85, 0.76]},
        ]
    catalog = source_image_catalog(results, scratch)
    assert catalog[0]["page"] == 1 and catalog[0]["figure_labels"] == ["Figure 1"]
    registry = freeze_selected_images(deck, results, scratch, settings)
    compiled, markdown = compile_deck(deck.model_dump(), results, {url: 1}, image_sources=registry)
    visual = compiled["slides"][0]["visual_data"]
    png = render_registered_page(registry, visual["id"])
    with Image.open(io.BytesIO(png)) as image:
        assert image.width > 400 and image.height <= 2200
    visual["sha256"] = hashlib.sha256(png).hexdigest()
    context = {
        "title": "Experiment",
        "markdown": markdown,
        "canonical_markdown": markdown,
        "stem": "source-page",
        "extras": {"deck": compiled, "slide_image_sources": registry},
        "citations": [url],
        "base_gates": [],
        "wants": ["pptx", "md"],
        "blocked": False,
        "template": "slides",
        "fail_on_quality": False,
        "generated_at": "2026-10-06T00:00:00Z",
    }
    preserved = [DeliveryFile("figures/" + visual["asset"], "png", "Original", "figure", png)]
    bundle = render_bundle(context, preserved)
    output = next(f for f in bundle.files if f.format == "pptx")
    assert output.status == "pass", output.issues
    ppt = Presentation(io.BytesIO(output.data))
    source_image = next(
        shape
        for slide in ppt.slides
        for shape in slide.shapes
        if shape.name.startswith("DR-image-")
    )
    assert source_image.image.blob == png
    text = "\n".join(s.text for slide in ppt.slides for s in slide.shapes if s.has_text_frame)
    selection = "用户选择区域" if region else "整页图"
    assert selection in text and "experiment.pdf" in text and "底层内容不可编辑" in text
    assert selection in markdown and "原文第 1 页" in markdown


async def test_source_registration_refuses_other_tasks_wrong_labels_and_tampering(
    tmp_path, settings
):
    settings, deck, results, scratch, url = await source_fixture(tmp_path, settings)
    with pytest.raises(ValueError, match="当前任务"):
        freeze_selected_images(deck, results, {}, settings)
    changed = deck.model_copy(deep=True)
    changed.slides[0].image.figure_label = "Figure 99"
    with pytest.raises(ValueError, match="图号"):
        freeze_selected_images(changed, results, scratch, settings)
    registry = freeze_selected_images(deck, results, scratch, settings)
    asset_id = deck.slides[0].image.asset_id
    broken = deepcopy(registry)
    next(iter(broken["documents"].values()))["sha256"] = "bad"
    with pytest.raises(ValueError, match="哈希"):
        render_registered_page(broken, asset_id)
    broken = deepcopy(registry)
    broken["images"][asset_id]["figure_label"] = "Figure 99"
    with pytest.raises(ValueError, match="图号|选择区域"):
        render_registered_page(broken, asset_id)
    with pytest.raises(ValueError, match="冻结来源"):
        compile_deck(deck.model_dump(), results, {url: 1}, image_sources=broken)


def test_model_cannot_choose_arbitrary_crops():
    with pytest.raises(ValueError):
        SlideDeck.model_validate(
            {
                "title": "X",
                "slides": [
                    {
                        "title": "X",
                        "image": {
                            "finding_id": "x",
                            "figure_label": "Figure 1",
                            "crop": [0, 0, 1, 1],
                        },
                    }
                ],
            }
        )


@pytest.mark.parametrize("bounds", [
    [-0.1, 0, 1, 1], [0, 0, 1.1, 1], [0.8, 0, 0.2, 1],
    [0, 0, 0, 1], [float("nan"), 0, 1, 1], [0, 0, float("inf"), 1],
])
def test_user_region_rejects_invalid_or_nonfinite_coordinates(bounds):
    from deep_research.workbench.attachments import ImageRegion

    with pytest.raises(ValueError):
        ImageRegion(page=1, figure_label="Figure 1", bounds=bounds)


async def test_user_selection_is_frozen_and_cannot_be_rebound_by_mutating_registry(
    tmp_path, settings,
):
    settings, deck, results, scratch, _ = await source_fixture(tmp_path, settings)
    scratch["attachments"][0]["image_regions"] = [
        {"page": 1, "figure_label": "Figure 1", "bounds": [0.1, 0.2, 0.9, 0.8]},
    ]
    registry = freeze_selected_images(deck, results, scratch, settings)
    asset_id = deck.slides[0].image.asset_id
    original = render_registered_page(registry, asset_id)
    scratch["attachments"][0]["image_regions"][0]["bounds"] = [0, 0, 1, 1]
    assert render_registered_page(registry, asset_id) == original
    broken = deepcopy(registry)
    broken["images"][asset_id]["bounds"] = [0, 0, 1, 1]
    with pytest.raises(ValueError, match="选择区域"):
        render_registered_page(broken, asset_id)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_user_region_matches_visible_pdf_coordinates_on_rotated_pages(rotation):
    import base64

    with fitz.open() as pdf:
        page = pdf.new_page(width=400, height=200)
        for rect, color in [
            ((0, 0, 200, 100), (1, 0, 0)), ((200, 0, 400, 100), (0, 1, 0)),
            ((0, 100, 200, 200), (0, 0, 1)), ((200, 100, 400, 200), (1, 1, 0)),
        ]:
            page.draw_rect(fitz.Rect(rect), color=color, fill=color)
        page.insert_text((10, 12), "Figure 1. Quadrants.")
        page.set_rotation(rotation)
        raw = pdf.tobytes()
        full = page.get_pixmap(alpha=False)
        expected = full.pixel(full.width // 4, full.height // 4)
    sha = hashlib.sha256(raw).hexdigest()
    page_identity = hashlib.sha256(f"{sha}:1".encode()).hexdigest()
    bounds = (0.0, 0.0, 0.5, 0.5)
    identity = f"{page_identity}:Figure 1:{bounds}".encode()
    asset = "source-" + hashlib.sha256(identity).hexdigest()[:20]
    registry = {
        "version": 2,
        "documents": {"doc": {
            "bytes": base64.b64encode(raw).decode(), "sha256": sha, "size": len(raw),
        }},
        "images": {asset: {
            "id": asset, "document_id": "doc", "page": 1, "figure_label": "Figure 1",
            "snapshot": "user_region", "selection_origin": "user", "bounds": list(bounds),
            "page_identity_sha256": page_identity,
        }},
    }
    with Image.open(io.BytesIO(render_registered_page(registry, asset))) as cropped:
        assert cropped.getpixel((cropped.width // 2, cropped.height // 2)) == expected
        assert max(cropped.size) <= 2200
