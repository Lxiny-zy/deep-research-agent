"""Actual PPTX objects, review inputs and bounded presentation contracts."""

from __future__ import annotations

import hashlib
import io
from copy import deepcopy
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData

from deep_research.models import Quantity, ResearchResult
from deep_research.workbench.delivery.pptx import paginate_deck, pptx_stats, render_pptx
from deep_research.workbench.delivery.pptx_visuals import visual_consistency
from deep_research.workbench.gates import slides_gate
from deep_research.workbench.slide_content import (
    SlideDeck,
    apply_brief,
    compile_deck,
    deck_to_markdown,
    slide_quality,
)
from deep_research.workbench.support import evidence_id
from tests.fakes import verified_finding


def numerical_deck(kind="table"):
    findings = []
    for name, number in [("A", 30.25), ("B", 31.5)]:
        finding = verified_finding(
            f"{name} PSNR 为 {number} dB", evidence_quote=f"{name} PSNR {number} dB"
        )
        finding.entity = name
        finding.quantity = Quantity(metric="PSNR", value=number, rendered=str(number), unit="dB")
        finding.verification.quantity_status = "verified"
        findings.append(finding)
    result = ResearchResult(sub_question="比较", findings=findings)
    raw = {
        "title": "比较",
        "slides": [
            {
                "title": "结果",
                "bullets": [],
                "notes": "逐项说明两种方法的实验结果与条件。",
                "citations": [1],
                "visual": {
                    "kind": kind,
                    "value_columns": ["p"],
                    "table": {
                        "id": "t1",
                        "title": "PSNR 比较",
                        "columns": [
                            {"key": "p", "label": "PSNR", "field": "quantity", "metric": "PSNR"}
                        ],
                        "rows": [
                            {"label": f.entity, "cells": {"p": [evidence_id(f)]}} for f in findings
                        ],
                    },
                },
            }
        ],
    }
    return raw, [result]


def test_empty_notes_and_pre_pagination_walls_are_not_hidden():
    raw = {
        "title": "汇报",
        "slides": [
            {"title": str(i), "bullets": ["密集正文" * 40] * 5, "notes": "（无备注）"}
            for i in range(8)
        ],
    }
    data = render_pptx(raw)
    assert pptx_stats(data)["with_notes"] == 1  # only the generated cover
    check = slides_gate(data, raw)
    assert check.status != "pass"
    assert check.metrics["original_pages"] == 8
    assert check.metrics["content_pages"] > 8
    assert check.metrics["overflow"] == 8


def test_continuations_partition_notes_without_loss_or_duplication():
    note = "先介绍背景。再解释方法。最后讨论局限。"
    raw = {
        "title": "汇报",
        "slides": [{"title": "长内容", "bullets": ["较长内容" * 35] * 3, "notes": note}],
    }
    pages = paginate_deck(raw)
    assert len(pages["slides"]) > 1
    assert "".join(p["notes"] for p in pages["slides"]) == note
    assert sum(p["notes"] == note for p in pages["slides"]) == 0
    assert paginate_deck(pages) == pages


def test_contract_overrides_model_invented_timing_and_detects_short_script():
    raw = SlideDeck(
        title="汇报",
        target_minutes=1,
        target_pages=99,
        slides=[{"title": "内容", "bullets": ["要点"], "notes": "很短的讲稿。"}],
    )
    deck = apply_brief(
        raw,
        SimpleNamespace(
            original_request="面向本科生，12 分钟，10 页", constraints=[], confirmed_choices={}
        ),
        "",
    )
    assert deck.target_minutes == 12 and deck.target_pages == 10
    issues, metrics = slide_quality(deck.model_dump())
    assert any("不足目标" in x for x in issues)
    assert metrics["timing_is_estimate"]
    assert apply_brief(raw, None, "不需要12分钟，5–8页").target_minutes is None
    assert apply_brief(raw, None, "12 分钟，语速每分钟 150 字").speaking_rate == 150


@pytest.mark.parametrize("kind", ["table", "bar", "line"])
def test_admitted_evidence_produces_real_editable_objects(kind):
    raw, results = numerical_deck(kind)
    source = deck_to_markdown(SlideDeck.model_validate(raw))
    assert "```evidence-table" in source
    compiled, markdown = compile_deck(raw, results, {"https://a.com": 1})
    assert "```evidence-table" not in markdown
    data = render_pptx(compiled)
    assert visual_consistency(data, compiled, markdown) == []
    pres = Presentation(io.BytesIO(data))
    shape = next(s for slide in pres.slides for s in slide.shapes if s.name.startswith("DR-"))
    if kind == "table":
        assert shape.has_table and "30.25" in shape.table.cell(1, 1).text
        shape.table.cell(1, 1).text = "999"
    else:
        assert shape.has_chart and list(shape.chart.series[0].values) == [30.25, 31.5]
        changed = CategoryChartData()
        changed.categories = ["A", "B"]
        changed.add_series("PSNR", [999, 31.5])
        shape.chart.replace_data(changed)
    altered = io.BytesIO()
    pres.save(altered)
    assert visual_consistency(altered.getvalue(), compiled, markdown)


def test_model_cannot_inject_rendered_data_or_use_unadmitted_findings():
    raw, results = numerical_deck("bar")
    raw["slides"][0]["visual_data"] = {"series": [{"values": [999]}]}
    compiled, _ = compile_deck(raw, results, {"https://a.com": 1})
    assert compiled["slides"][0]["visual_data"]["series"][0]["values"] == [30.25, 31.5]
    with pytest.raises(ValueError, match="不存在|未通过"):
        compile_deck(raw, [], {"https://a.com": 1})


def test_missing_or_uncertain_chart_values_are_not_zero_filled():
    raw, results = numerical_deck("bar")
    results[0].findings[0].quantity.uncertainty = 0.5
    with pytest.raises(ValueError, match="不确定度"):
        compile_deck(raw, results, {"https://a.com": 1})


def test_images_require_frozen_hash_and_preserve_pixels_not_fake_editability():
    image = io.BytesIO()
    Image.new("RGB", (180, 100), "white").save(image, format="PNG")
    blob = image.getvalue()
    visual = {
        "kind": "image",
        "id": "source",
        "asset": "fig.png",
        "caption": "原图 1 [1]",
        "citations": [1],
        "sha256": hashlib.sha256(blob).hexdigest(),
    }
    deck = {
        "title": "方法",
        "slides": [
            {"title": "原图", "bullets": [], "notes": "说明图中步骤。", "visual_data": visual}
        ],
    }
    pptx = render_pptx(deck, images={"fig.png": blob})
    assert visual_consistency(pptx, deck) == []
    pres = Presentation(io.BytesIO(pptx))
    text = "\n".join(s.text for slide in pres.slides for s in slide.shapes if s.has_text_frame)
    assert "底层内容不可编辑" in text
    tampered = deepcopy(deck)
    tampered["slides"][0]["visual_data"]["sha256"] = "bad"
    with pytest.raises(ValueError, match="哈希"):
        render_pptx(tampered, images={"fig.png": blob})
    with pytest.raises(ValueError, match="冻结"):
        render_pptx(deck)


@pytest.mark.parametrize("kind", ["table", "bar", "line"])
async def test_writer_table_projection_survives_frozen_delivery(kind, settings):
    from deep_research.agents.base import Blackboard, RunContext
    from deep_research.models import Report
    from deep_research.observability import Tracer
    from deep_research.workbench.delivery_render import render_bundle
    from deep_research.workbench.tables import render_specs
    from deep_research.workbench.templates import SLIDES
    from deep_research.workbench.writers import SlideWriter
    from tests.fakes import FakeLLM, FakeSearch

    raw, results = numerical_deck(kind)

    class VisualWriter(FakeLLM):
        async def parse(self, system, user, schema, **kwargs):
            assert schema is SlideDeck
            assert "evidence-table" not in system  # one JSON output contract
            return SlideDeck.model_validate(raw)

    bb = Blackboard(query="比较两个方法", results=results)
    ctx = RunContext(
        llm=VisualWriter(), search_tool=FakeSearch(), tracer=Tracer(), settings=settings
    )
    writer = SlideWriter()
    draft = await writer.write(bb, ctx, SLIDES, None, "固定测试材料")
    body, record = render_specs(draft, results, {"https://a.com": 1})
    assert record["specs"] and not record["errors"]
    report = Report(query=bb.query, markdown=body, citations=["https://a.com"])
    extras = writer.postprocess(bb, report, SLIDES)
    assert extras.get("deck") and not extras.get("structured_output_issue")
    # Isolate the frozen-format boundary. Evidence correctness is exercised by
    # the existing table/prose suites, not declared from this fixed fixture.
    context = {
        "title": "比较",
        "markdown": body,
        "canonical_markdown": body,
        "stem": "slides",
        "extras": extras,
        "citations": ["https://a.com"],
        "base_gates": [],
        "wants": ["pptx", "md"],
        "blocked": False,
        "template": "slides",
        "fail_on_quality": False,
        "generated_at": "2026-10-06T00:00:00Z",
    }
    bundle = render_bundle(context, [])
    file = next(f for f in bundle.files if f.format == "pptx")
    assert file.status == "pass", file.issues
    assert visual_consistency(file.data, extras["deck"], body) == []
    assert "30.25" in next(f.data.decode() for f in bundle.files if f.format == "md")
    changed = {**context, "markdown": body.replace("30.25", "99.25")}
    failed = render_bundle(changed, [])
    assert next(f for f in failed.files if f.format == "pptx").status == "fail"


def test_partition_keeps_repeated_punctuation_and_spacing():
    text = "第一段！！  Second paragraph!\n第三段。\n"
    deck = {
        "title": "检查",
        "slides": [
            {
                "title": "内容",
                "bullets": ["密集内容" * 35] * 3,
                "notes": text,
            }
        ],
    }
    assert "".join(s["notes"] for s in paginate_deck(deck)["slides"]) == text


def test_native_chart_workbook_preserves_literal_source_labels():
    import zipfile

    import openpyxl

    raw, results = numerical_deck("bar")
    compiled, _ = compile_deck(raw, results, {"https://a.com": 1})
    visual = compiled["slides"][0]["visual_data"]
    visual["categories"] = ['=HYPERLINK("https://invalid","x")', "https://invalid"]
    visual["series"][0]["name"] = "=1+1"
    data = render_pptx(compiled)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        workbook_name = next(n for n in archive.namelist() if n.endswith(".xlsx"))
        workbook = openpyxl.load_workbook(io.BytesIO(archive.read(workbook_name)))
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["B1"].data_type == "s"
    assert workbook.active["A3"].hyperlink is None
    assert workbook.active["B1"].value == "=1+1"


def test_bibliography_links_do_not_break_native_table_consistency():
    from deep_research.bibliography import build_bibliography, project_citations
    from deep_research.workbench.delivery_render import render_bundle

    raw, results = numerical_deck("bar")
    compiled, markdown = compile_deck(raw, results, {"https://a.com": 1})
    bibliography = build_bibliography(markdown, ["https://a.com"], results[0].findings)
    shown = project_citations(markdown, bibliography)
    context = {
        "title": "比较",
        "markdown": shown,
        "canonical_markdown": markdown,
        "stem": "bib",
        "extras": {"deck": compiled},
        "citations": ["https://a.com"],
        "base_gates": [],
        "bibliography": bibliography.model_dump(mode="json"),
        "wants": ["pptx", "md"],
        "blocked": False,
        "template": "slides",
        "fail_on_quality": False,
        "generated_at": "2026-10-06T00:00:00Z",
    }
    bundle = render_bundle(context, [])
    output = next(f for f in bundle.files if f.format == "pptx")
    assert output.status == "pass", output.issues
