"""Actual input masks, frozen selection counts, and workbook result navigation."""

import io
from copy import deepcopy
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from deep_research.workbench.analysis import analyse, fallback_report
from deep_research.workbench.analysis_lineage import attach_design, build_lineage, missing_rows
from deep_research.workbench.analysis_review import correlation_interpretation_issues
from deep_research.workbench.publish import _stats_xlsx

CSV = """id,group,x,y,unused
001,A,1.1,5.8,meta
002,A,2.2,?,meta
003,A,NA,9.2,meta
004,B,3.3,10.4,meta
005,B,4.9,12,meta
006,,5.1,14.2,meta
007,B,6.8,15.6,meta
008,A,7.4,19,
"""
SCOPE = {
    "measures": ["x", "y"],
    "groups": ["group"],
    "background": ["id", "unused"],
    "subject_columns": ["id"],
    "comparison": "independent",
}


@pytest.fixture
def result():
    return analyse(CSV, "按group比较x和y", scope=SCOPE)


def test_missing_positions_and_statistic_pools_are_actual_input_subsets(result):
    assert list(missing_rows(result.lineage, expected_hash=result.input_sha256)) == [
        (2, ["y"]),
        (3, ["x"]),
        (6, ["group"]),
        (8, ["unused"]),
    ]
    rules = result.lineage["rules"]
    description = next(r for r in rules if r["uses"] == ["描述统计：x"])
    comparison = next(r for r in rules if r["required_columns"] == ["x", "group"])
    correlation = next(r for r in rules if r["required_columns"] == ["x", "y"] and not r["where"])
    assert description["included_rows"] == 7
    assert comparison["included_rows"] == 6 and comparison["excluded_missing_rows"] == 2
    assert correlation["included_rows"] == 6
    summary = attach_design("## 分析计划\n比较变量。", result.design, result.lineage)
    assert "独立性仍依赖实验设计" in summary and "未自动删除" in summary
    assert "对象标识列：id" in summary
    assert attach_design(summary, result.design, result.lineage) == summary


def test_frozen_lineage_and_design_are_not_recreated_under_new_parser(result):
    frozen = result.snapshot()
    assert frozen["version"] == 6
    restored = analyse(CSV, frozen=frozen)
    assert restored.snapshot() == frozen
    legacy_result = deepcopy(result)
    legacy_result.design = legacy_result.lineage = None
    legacy = legacy_result.snapshot()
    restored_legacy = analyse(CSV, frozen=legacy)
    assert restored_legacy.snapshot() == legacy
    workbook = load_workbook(io.BytesIO(_stats_xlsx(restored_legacy)))
    assert workbook["缺失记录"]["A2"].value == "未记录"


def test_workbook_links_report_claims_to_existing_statistic_ids_without_subject_values(result):
    mean = next(r["mean"] for r in result.describe if r["variable"] == "x")
    result.report_markdown = f"## 统计结果\nx 的均值为 {mean}。"
    workbook = load_workbook(io.BytesIO(_stats_xlsx(result)))
    assert list(workbook["缺失记录"].values)[1:] == [
        (2, "y"),
        (3, "x"),
        (6, "group"),
        (8, "unused"),
    ]
    location = workbook["正文结果定位"]
    assert location["E2"].hyperlink and "统计记录索引" in location["E2"].hyperlink.location
    fact_id = location["E2"].value
    record = next(row for row in workbook["统计记录索引"].values if row[0] == fact_id)
    assert record[2] == "x" and record[7] == "mean" and record[8] == mean
    assert record[9] == result.input_sha256
    assert not any(
        cell.value in {"001", "002", "003", "008"}
        for sheet in workbook
        for row in sheet
        for cell in row
        if isinstance(cell.value, str)
    )
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


def test_changed_lineage_hash_is_not_silently_accepted(result):
    result.lineage["source_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="冻结输入"):
        _stats_xlsx(result)


def test_missing_index_size_is_bounded(monkeypatch):
    import deep_research.workbench.analysis_lineage as module

    monkeypatch.setattr(module, "MAX_BITMAP_BYTES", 1)
    with pytest.raises(ValueError, match="容量"):
        build_lineage(pd.DataFrame({"x": [None] * 100}), SimpleNamespace())


def test_repeated_subject_design_remains_descriptive_and_explicitly_unsupported():
    csv = "subject,time,x\na,1,1.1\na,2,2.7\nb,1,3.2\nb,2,4.6"
    value = analyse(
        csv,
        scope={
            "measures": ["x"],
            "groups": [],
            "background": ["subject", "time"],
            "subject_columns": ["subject"],
        },
    )
    assert value.design["comparison"] == "repeated_unsupported"
    assert value.tests == [] and "不支持重复测量推断" in fallback_report(value)


@pytest.mark.parametrize(
    "statement,bad",
    [
        ("x 与 y 呈负相关。", True),
        ("x 与 y 呈正相关。", False),
        ("A组的 x 与 y 呈负相关。", False),
        ("A组的 x 与 y 呈正相关。", True),
        ("B组的 x 与 y 呈显著正相关。", True),
        ("x 导致 y 增大。", True),
        ("x 与 y 相关不能证明因果关系。", False),
        ("假设 x 可能导致 y，仍需要进一步实验验证。", False),
    ],
)
def test_interpretation_checks_use_the_specific_correlation_scope(statement, bad):
    ledger = {
        "correlations": [
            {"a": "x", "b": "y", "r": 0.8, "p_value": 0.01},
            {"a": "x", "b": "y", "r": -0.7, "p_value": 0.02, "group_column": "kind", "group": "A"},
            {"a": "x", "b": "y", "r": 0.2, "p_value": 0.4, "group_column": "kind", "group": "B"},
        ]
    }
    assert bool(correlation_interpretation_issues(statement, ledger)) is bad


async def test_actual_workbench_freezes_design_and_exports_result_locator(settings):
    from deep_research.workbench.publish import build_bundle
    from tests.test_workbench import _run

    query = "按method比较psnr。\nmethod,psnr\nA,30\nA,30.5\nA,31\nB,33\nB,33.2\nB,34"
    report, detail, _ = await _run("dataAnalysis", query, "## 结论\n均值为99.99。", settings)
    assert "分析设计与记录范围" in report.markdown
    ledger = detail.orchestration.checkpoint["scratch"]["analysis"]
    assert ledger["version"] == 6 and ledger["lineage"]["source_sha256"] == ledger["input_sha256"]
    bundle = build_bundle(detail)
    file = next(file for file in bundle.files if file.format == "xlsx")
    workbook = load_workbook(io.BytesIO(file.data))
    assert workbook["统计记录索引"].max_row > 2
    locations = workbook["正文结果定位"]
    assert any(cell.hyperlink for row in locations for cell in row)
    assert bundle.render_context["statistics"]["lineage"] == ledger["lineage"]
